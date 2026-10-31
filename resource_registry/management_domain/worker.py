"""Prepare verified files outside transactions and fence stale worker results."""

from __future__ import annotations

import hashlib
import logging
from datetime import timedelta
from pathlib import Path
from tempfile import TemporaryDirectory
from uuid import UUID, uuid4

from django.db import transaction
from django.utils import timezone

# The existing Excel parser has no installed type stubs.
from openpyxl import load_workbook  # type: ignore[import-untyped]

from resource_registry.cleanup.inventory import register_object
from resource_registry.management_domain.changes import DatasetChangeService, JsonObject
from resource_registry.management_domain.comparison import compare_dataset_schemas
from resource_registry.management_domain.errors import DatasetManagementError
from resource_registry.models import Dataset, DatasetChange, DatasetVersion, StagedResourceIntake
from resource_registry.services.canonical_dataset_builder import CanonicalDatasetBuilder
from resource_registry.services.dataset_context_builder import DatasetContextBuilder
from resource_registry.storage.factory import DataRegistryStorageFactory

logger = logging.getLogger(__name__)


class DatasetChangeWorker:
    """Prepare private candidates using the existing canonicalization and storage services."""

    def run_once(self, *, batch_size: int = 5, change_id: UUID | None = None) -> int:
        """Process a bounded batch, recovering abandoned preparation for explicit retry.

        Args:
            batch_size: Maximum proposals claimed by this invocation.
            change_id: Optional exact proposal to process for an automatic wake-up.

        Returns:
            Number of proposals claimed, including failed preparations.
        """
        candidates = DatasetChange.objects.all()
        if change_id is not None:
            candidates = candidates.filter(pk=change_id)
        candidates.filter(
            status="PROCESSING",
            started_at__lt=timezone.now() - timedelta(hours=1),
        ).update(
            status="FAILED",
            error_code="worker_interrupted",
            error_message="Preparation was interrupted. Retry this update.",
        )
        claimed = 0
        for _ in range(max(0, batch_size)):
            with transaction.atomic():
                change = (
                    candidates.select_for_update(skip_locked=True)
                    .filter(
                        status="QUEUED",
                        namespace__workspace__organization__datasetmanagementconfiguration__enabled=True,
                    )
                    .order_by("created_at")
                    .first()
                )
                if change is None:
                    change = (
                        candidates.select_for_update(skip_locked=True)
                        .filter(
                            status="QUEUED",
                            namespace__organization__datasetmanagementconfiguration__enabled=True,
                        )
                        .order_by("created_at")
                        .first()
                    )
                if change is None:
                    break
                change.status = "PROCESSING"
                change.started_at = timezone.now()
                change.save(update_fields=["status", "started_at", "updated_at"])
            claimed += 1
            self._prepare_claimed_change(change)
        return claimed

    def _prepare_claimed_change(self, change: DatasetChange) -> None:
        try:
            namespace = change.namespace
            if namespace.organization_id is None and namespace.workspace is None:
                raise DatasetManagementError(
                    "invalid_destination", "The destination no longer exists."
                )
            organization = (
                namespace.organization
                if namespace.organization_id
                else namespace.workspace.organization if namespace.workspace else None
            )
            service = DatasetChangeService(change.submitted_by, organization)
            service.authorization.require_enabled()
            service.validate_external_change(change, change.external_source)
            service.validate_targets(
                namespace, change.specification, external_source=change.external_source
            )
            prepared = self._prepare_file(change)
            items = prepared["items"]
            if not items:
                status = "READY"
            elif all(
                item["unchanged"] and not item.get("mapping_changed") for item in items
            ) and not change.specification.get("archive"):
                status = "NO_CHANGES"
            elif (change.external_source_id and change.specification.get("archive")) or any(
                item["difference"]["requires_manager"] or item.get("historical_version_id")
                for item in items
            ):
                status = "NEEDS_MANAGER"
            else:
                status = "READY"
            if change.external_source_id:
                service.verify_external_versions(change, prepared)
                service.complete_external_preparation(
                    change.id, expected_revision=change.revision, prepared=prepared, status=status
                )
                return
            DatasetChange.objects.filter(
                pk=change.pk, revision=change.revision, status="PROCESSING"
            ).update(
                prepared=prepared,
                status=status,
                error_code="",
                error_message="",
                completed_at=timezone.now() if status == "NO_CHANGES" else None,
                updated_at=timezone.now(),
            )
        except Exception as error:
            # The durable worker boundary must record failures, including third-party parsers/I/O.
            code = error.code if isinstance(error, DatasetManagementError) else "preparation_failed"
            message = (
                str(error)
                if isinstance(error, DatasetManagementError)
                else "File preparation failed. Check the file and retry."
            )
            logger.warning(
                "Dataset change preparation failed: change=%s code=%s exception=%s",
                change.id,
                code,
                type(error).__name__,
            )
            DatasetChange.objects.filter(
                pk=change.pk, revision=change.revision, status="PROCESSING"
            ).update(
                status="FAILED",
                error_code=code,
                error_message=message[:255],
                updated_at=timezone.now(),
            )

    def _prepare_file(self, change: DatasetChange) -> JsonObject:
        source = (
            DatasetChangeService.record_intake_source_use(change.intake_id)
            if change.intake_id
            else change.upload
        )
        if source is None or source.status not in {
            "READY_FOR_PUBLICATION",
            "PUBLISHED",
            "COMPLETED",
        }:
            raise DatasetManagementError(
                "source_unavailable", "The verified source is unavailable."
            )
        DatasetChangeService.validate_source_mapping(
            change.specification, file_format=source.resolved_file_format
        )
        provider = DataRegistryStorageFactory().create_provider(source.storage_provider)
        filename = Path(source.normalized_filename).name
        source_format = source.resolved_file_format.lower()
        prefix = f"dataset-changes/{change.namespace_id}/{change.id}/{change.revision}/{uuid4()}"
        with TemporaryDirectory(prefix="dataset-change-") as directory:
            source_path = Path(directory) / filename
            provider.download_to_file(provider.get_uri(source.storage_object_key), str(source_path))
            checksum = self._checksum(source_path)
            expected_checksum = source.verified_checksum
            if not expected_checksum and isinstance(source, StagedResourceIntake):
                expected_checksum = source.inspection_json.get("content_hash")
            if source.actual_size_bytes != source_path.stat().st_size or (
                expected_checksum and checksum != expected_checksum
            ):
                raise DatasetManagementError(
                    "source_changed",
                    "The file no longer matches its verified upload. Upload it again.",
                )
            stored_source = provider.put(str(source_path), f"{prefix}/source/{filename}")

            register_object(
                provider=source.storage_provider,
                uri=stored_source.uri,
                owner_kind="CHANGE",
                owner_id=change.id,
                revision=change.revision,
                verified=True,
                checksum=checksum,
                etag=stored_source.etag or "",
            )
            is_workbook = source_format == "xlsx"
            available = self._list_sheets(source_path) if is_workbook else []
            DatasetChange.objects.filter(
                pk=change.pk, revision=change.revision, status="PROCESSING"
            ).update(
                prepared={
                    "available_sheets": available,
                    "suggested_sheets": self._suggest_sheets(change, available),
                    "items": [],
                    "filename": filename,
                    "is_workbook": is_workbook,
                },
            )
            selected = (
                self._select_sheets(change, available)
                if is_workbook
                else [
                    {
                        "dataset_id": change.specification.get("dataset_id"),
                        "name": change.specification.get("name") or Path(filename).stem,
                    }
                ]
            )
            if not is_workbook and change.specification.get("workbook_id"):
                raise DatasetManagementError(
                    "workbook_required", "Upload the complete Excel workbook."
                )
            items = []
            for index, entry in enumerate(selected):
                if entry.get("sheet_name") and entry["sheet_name"] not in available:
                    raise DatasetManagementError(
                        "sheet_missing", "A selected worksheet is missing. Revise the mapping."
                    )
                with CanonicalDatasetBuilder().build(
                    source_path=str(source_path),
                    source_format=source_format,
                    sheet_name=entry.get("sheet_name"),
                ) as canonical:
                    context = DatasetContextBuilder().build_context(
                        canonical.canonical_path, preview_limit=0
                    )
                    if not context.get("schema"):
                        raise DatasetManagementError(
                            "invalid_sheet",
                            "A selected worksheet has no usable columns. Exclude it and recalculate.",
                        )
                    context.pop("preview", None)
                    context["source"]["path"] = stored_source.uri
                    context["source"]["file_name"] = filename
                    context["source"]["sheet_name"] = entry.get("sheet_name")
                    canonical_hash = self._checksum(Path(canonical.canonical_path))
                    artifact = provider.put(
                        canonical.canonical_path, f"{prefix}/sheets/{index}.parquet"
                    )
                    register_object(
                        provider=source.storage_provider,
                        uri=artifact.uri,
                        owner_kind="CHANGE",
                        owner_id=change.id,
                        revision=change.revision,
                        verified=True,
                        checksum=canonical_hash,
                        etag=artifact.etag or "",
                    )
                    dataset = (
                        Dataset.objects.select_related("current_version")
                        .filter(pk=str(entry["dataset_id"]))
                        .first()
                        if entry.get("dataset_id")
                        else None
                    )
                    current = dataset.current_version if dataset else None
                    old_schema = (
                        current.schemas.order_by("-created_at").first() if current else None
                    )
                    old_profile = (
                        current.profiles.order_by("-computed_at").first() if current else None
                    )
                    difference = compare_dataset_schemas(
                        old_schema.schema_json if old_schema else [], context["schema"]
                    )
                    if current is not None and old_schema is None:
                        difference["requires_manager"] = True
                    difference.update(
                        previous_rows=(
                            old_profile.metrics_json.get("row_count") if old_profile else None
                        ),
                        new_rows=context["summary"]["row_count"],
                    )
                    match = (
                        DatasetVersion.objects.filter(
                            dataset=dataset, status="READY", content_hash=canonical_hash
                        ).first()
                        if dataset
                        else None
                    )
                    if not match and dataset:
                        match = DatasetVersion.objects.filter(
                            dataset=dataset,
                            status="READY",
                            artifacts__checksum=canonical_hash,
                            artifacts__metadata_json__is_canonical=True,
                        ).first()
                    items.append(
                        {
                            "dataset_id": str(dataset.id) if dataset else None,
                            "name": entry.get("name")
                            or (
                                dataset.name
                                if dataset
                                else entry.get("sheet_name") or Path(filename).stem
                            ),
                            "sheet_name": entry.get("sheet_name"),
                            "context": context,
                            "content_hash": canonical_hash,
                            "checksum": canonical_hash,
                            "storage_etag": artifact.etag,
                            "uri": artifact.uri,
                            "size_bytes": artifact.size_bytes,
                            "provider": source.storage_provider,
                            "unchanged": bool(match and current and match.id == current.id),
                            "mapping_changed": bool(
                                is_workbook
                                and dataset
                                and (
                                    not dataset.workbook_id
                                    or dataset.workbook_sheet != entry.get("sheet_name")
                                )
                            ),
                            "historical_version_id": (
                                str(match.id)
                                if match and (not current or match.id != current.id)
                                else None
                            ),
                            "difference": difference,
                        }
                    )
        return {
            "items": items,
            "source_uri": stored_source.uri,
            "source_hash": checksum,
            "filename": filename,
            "is_workbook": is_workbook,
            "available_sheets": available,
            "suggested_sheets": self._suggest_sheets(change, available),
        }

    @staticmethod
    def _checksum(path: Path) -> str:
        with path.open("rb") as stream:
            return hashlib.file_digest(stream, "sha256").hexdigest()

    @staticmethod
    def _list_sheets(path: Path) -> list[str]:
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            return list(workbook.sheetnames)
        finally:
            workbook.close()

    @staticmethod
    def _select_sheets(change: DatasetChange, available: list[str]) -> list[JsonObject]:
        return [
            entry for entry in change.specification.get("sheets", []) if entry.get("selected", True)
        ]

    @staticmethod
    def _suggest_sheets(change: DatasetChange, available: list[str]) -> list[JsonObject]:
        datasets = (
            {
                dataset.workbook_sheet: dataset
                for dataset in Dataset.objects.filter(
                    workbook_id=change.workbook_id, archived_at__isnull=True
                ).exclude(status="ARCHIVED")
            }
            if change.workbook_id
            else {}
        )
        explicit = {entry["sheet_name"]: entry for entry in change.specification.get("sheets", [])}
        return [
            explicit.get(name)
            or {
                "sheet_name": name,
                "dataset_id": str(datasets[name].id) if name in datasets else None,
                "expected_revision": datasets[name].revision if name in datasets else 0,
                "selected": name in datasets,
            }
            for name in available
        ]
