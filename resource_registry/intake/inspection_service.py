"""Run durable inspection jobs for private staged resource intakes."""

from __future__ import annotations

import hashlib
import uuid
from datetime import timedelta
from pathlib import Path
from tempfile import NamedTemporaryFile
from typing import Any, Protocol

from django.conf import settings
from django.db import transaction
from django.utils import timezone
from openpyxl import load_workbook  # type: ignore[import-untyped]  # Openpyxl has no bundled stubs.

from resource_registry.cleanup.inventory import register_intake_object
from resource_registry.models import (
    StagedResourceInspectionJob,
    StagedResourceInspectionJobStatus,
    StagedResourceIntake,
    StagedResourceIntakeStatus,
)
from resource_registry.services.dataset_context_builder import DatasetContextBuilder
from resource_registry.services.json_dataset import JsonDatasetValidationError
from resource_registry.services.parquet_dataset import (
    ParquetDatasetValidationError,
    validate_parquet_dataset,
)
from resource_registry.storage.factory import DataRegistryStorageFactory
from resource_registry.uploads.contracts import SupportedUploadFormat
from resource_registry.uploads.exceptions import UploadApplicationError

_MAX_INSPECTION_ATTEMPTS = 3


class InspectionFunction(Protocol):
    """Define the compatibility seam for one staged-file inspection."""

    def __call__(self, *, intake: StagedResourceIntake) -> dict[str, object]:
        """Inspect a staged intake and return safe metadata."""


class StagedResourceIntakeInspectionService:
    """Execute inspection jobs without taking ownership of intake publication."""

    def __init__(
        self,
        *,
        storage_factory: DataRegistryStorageFactory,
        context_builder: DatasetContextBuilder,
    ) -> None:
        """Initialize immutable dependencies used to inspect staged source files.

        Args:
            storage_factory: Creates the provider that owns each staged object.
            context_builder: Extracts safe, non-persistent file summaries.
        """
        self._storage_factory = storage_factory
        self._context_builder = context_builder

    def run_jobs(
        self,
        *,
        batch_size: int,
        inspect_uploaded_file: InspectionFunction | None = None,
        intake_id: uuid.UUID | None = None,
    ) -> int:
        """Claim and execute a bounded batch of private intake inspections.

        Args:
            batch_size: Maximum number of queued jobs to claim.
            inspect_uploaded_file: Optional compatibility hook for inspecting one file.
            intake_id: Optional exact intake for an authorized automatic wake-up.

        Returns:
            Number of inspections completed successfully.
        """
        completed = 0
        self._recover_abandoned_jobs(intake_id=intake_id)
        inspection_function = inspect_uploaded_file or self.inspect_uploaded_file
        candidates = StagedResourceInspectionJob.objects.all()
        if intake_id is not None:
            candidates = candidates.filter(intake_id=intake_id)
        for _ in range(max(int(batch_size), 1)):
            with transaction.atomic():
                job = (
                    candidates.select_for_update(skip_locked=True)
                    .select_related("intake")
                    .filter(status=StagedResourceInspectionJobStatus.QUEUED)
                    .order_by("created_at", "id")
                    .first()
                )
                if job is None:
                    break
                if self._has_intake_expired(job.intake):
                    job.status = StagedResourceInspectionJobStatus.FAILED
                    job.error_code = "INTAKE_EXPIRED"
                    job.completed_at = timezone.now()
                    job.save(update_fields=["status", "error_code", "completed_at", "updated_at"])
                    job.intake.status = StagedResourceIntakeStatus.EXPIRED
                    job.intake.cleanup_required = True
                    job.intake.save(update_fields=["status", "cleanup_required", "updated_at"])
                    continue
                job.status = StagedResourceInspectionJobStatus.RUNNING
                job.started_at = timezone.now()
                job.save(update_fields=["status", "started_at", "updated_at"])
                job.intake.status = StagedResourceIntakeStatus.ANALYZING
                job.intake.save(update_fields=["status", "updated_at"])
                claimed_intake_id = job.intake_id
            try:
                intake = StagedResourceIntake.objects.get(pk=claimed_intake_id)
                inspection = inspection_function(intake=intake)
            except (JsonDatasetValidationError, ParquetDatasetValidationError) as error:
                self._record_terminal_inspection_failure(
                    job_id=job.pk,
                    intake_id=claimed_intake_id,
                    error_code=str(error),
                )
                break
            except Exception:
                self._record_inspection_failure(job_id=job.pk, intake_id=claimed_intake_id)
                break
            if self._persist_inspection_success(
                job_id=job.pk,
                intake_id=claimed_intake_id,
                inspection=inspection,
            ):
                completed += 1
        return completed

    def inspect_uploaded_file(self, *, intake: StagedResourceIntake) -> dict[str, object]:
        """Inspect one uploaded file and return only safe workbook metadata.

        Args:
            intake: Intake whose server-selected object is inspected.

        Returns:
            Content digest and per-sheet summaries.

        Raises:
            UploadApplicationError: If a declared checksum does not match the object.
        """
        provider = self._storage_factory.create_provider(intake.storage_provider)
        source_uri = provider.get_uri(intake.storage_object_key)
        with NamedTemporaryFile(
            suffix=Path(intake.normalized_filename).suffix, delete=True
        ) as handle:
            provider.download_to_file(source_uri, handle.name)
            content_hash = self.calculate_sha256_file(Path(handle.name))
            if intake.declared_checksum and content_hash != intake.declared_checksum:
                raise UploadApplicationError("INTAKE_CHECKSUM_MISMATCH", 409)
            if intake.resolved_file_format == SupportedUploadFormat.PARQUET.value:
                validate_parquet_dataset(handle.name)
            if intake.resolved_file_format in {
                SupportedUploadFormat.CSV.value,
                SupportedUploadFormat.JSON.value,
                SupportedUploadFormat.JSONL.value,
                SupportedUploadFormat.PARQUET.value,
            }:
                context = self._context_builder.build_context(handle.name, preview_limit=0)
                return {
                    "content_hash": content_hash,
                    "sheets": [
                        self._safe_sheet_summary(name=intake.resolved_file_format, context=context)
                    ],
                }
            workbook = load_workbook(handle.name, read_only=True, data_only=True)
            try:
                sheets: list[dict[str, object]] = []
                for sheet_name in workbook.sheetnames:
                    try:
                        context = self._context_builder.build_context(
                            handle.name,
                            preview_limit=0,
                            sheet_name=sheet_name,
                        )
                        sheets.append(self._safe_sheet_summary(name=sheet_name, context=context))
                    except Exception:
                        sheets.append({"name": sheet_name, "status": "REJECTED"})
                return {"content_hash": content_hash, "sheets": sheets}
            finally:
                workbook.close()

    def _record_inspection_failure(self, *, job_id: uuid.UUID, intake_id: uuid.UUID) -> bool:
        """Persist retry, terminal failure, or expiry after one inspection error.

        Args:
            job_id: Primary key of the claimed inspection job.
            intake_id: Primary key of the associated intake.

        Returns:
            Whether job processing should stop for this batch.
        """
        with transaction.atomic():
            failed_job = StagedResourceInspectionJob.objects.select_for_update().get(pk=job_id)
            failed_intake = StagedResourceIntake.objects.select_for_update().get(pk=intake_id)
            if (
                failed_intake.status == StagedResourceIntakeStatus.EXPIRED
                or self._has_intake_expired(failed_intake)
            ):
                failed_job.status = StagedResourceInspectionJobStatus.FAILED
                failed_job.error_code = "INTAKE_EXPIRED"
                failed_job.completed_at = timezone.now()
                failed_job.save(
                    update_fields=["status", "error_code", "completed_at", "updated_at"]
                )
                if failed_intake.status != StagedResourceIntakeStatus.EXPIRED:
                    failed_intake.status = StagedResourceIntakeStatus.EXPIRED
                    failed_intake.cleanup_required = True
                    failed_intake.failure_code = "INTAKE_EXPIRED"
                    failed_intake.save(
                        update_fields=["status", "cleanup_required", "failure_code", "updated_at"]
                    )
                return True
            failed_job.retry_count += 1
            if failed_job.retry_count < _MAX_INSPECTION_ATTEMPTS:
                failed_job.status = StagedResourceInspectionJobStatus.QUEUED
                failed_job.error_code = "INTAKE_INSPECTION_RETRYABLE"
                failed_job.started_at = None
                failed_job.save(
                    update_fields=[
                        "status",
                        "error_code",
                        "retry_count",
                        "started_at",
                        "updated_at",
                    ]
                )
                failed_intake.status = StagedResourceIntakeStatus.COMPLETED
                failed_intake.save(update_fields=["status", "updated_at"])
            else:
                failed_job.status = StagedResourceInspectionJobStatus.FAILED
                failed_job.error_code = "INTAKE_INSPECTION_FAILED"
                failed_job.completed_at = timezone.now()
                failed_job.save(
                    update_fields=[
                        "status",
                        "error_code",
                        "retry_count",
                        "completed_at",
                        "updated_at",
                    ]
                )
                failed_intake.status = StagedResourceIntakeStatus.FAILED
                failed_intake.failure_code = "INTAKE_INSPECTION_FAILED"
                failed_intake.save(update_fields=["status", "failure_code", "updated_at"])
        return False

    def _record_terminal_inspection_failure(
        self,
        *,
        job_id: uuid.UUID,
        intake_id: uuid.UUID,
        error_code: str,
    ) -> None:
        """Persist a non-retryable inspection failure and its safe error code.

        Args:
            job_id: Primary key of the claimed inspection job.
            intake_id: Primary key of the associated intake.
            error_code: Stable, safe validation code suitable for operator visibility.
        """
        with transaction.atomic():
            failed_job = StagedResourceInspectionJob.objects.select_for_update().get(pk=job_id)
            failed_intake = StagedResourceIntake.objects.select_for_update().get(pk=intake_id)
            if (
                failed_intake.status == StagedResourceIntakeStatus.EXPIRED
                or self._has_intake_expired(failed_intake)
            ):
                failed_job.status = StagedResourceInspectionJobStatus.FAILED
                failed_job.error_code = "INTAKE_EXPIRED"
                failed_job.completed_at = timezone.now()
                failed_job.save(
                    update_fields=["status", "error_code", "completed_at", "updated_at"]
                )
                if failed_intake.status != StagedResourceIntakeStatus.EXPIRED:
                    failed_intake.status = StagedResourceIntakeStatus.EXPIRED
                    failed_intake.cleanup_required = True
                    failed_intake.failure_code = "INTAKE_EXPIRED"
                    failed_intake.save(
                        update_fields=["status", "cleanup_required", "failure_code", "updated_at"]
                    )
                return
            failed_job.status = StagedResourceInspectionJobStatus.FAILED
            failed_job.error_code = error_code
            failed_job.completed_at = timezone.now()
            failed_job.save(update_fields=["status", "error_code", "completed_at", "updated_at"])
            failed_intake.status = StagedResourceIntakeStatus.FAILED
            failed_intake.failure_code = error_code
            failed_intake.save(update_fields=["status", "failure_code", "updated_at"])

    def _persist_inspection_success(
        self,
        *,
        job_id: uuid.UUID,
        intake_id: uuid.UUID,
        inspection: dict[str, object],
    ) -> bool:
        """Persist a completed inspection only while its intake remains active.

        Args:
            job_id: Primary key of the claimed inspection job.
            intake_id: Primary key of the associated intake.
            inspection: Safe inspection result to persist.

        Returns:
            Whether the inspection completed successfully.
        """
        with transaction.atomic():
            locked_job = StagedResourceInspectionJob.objects.select_for_update().get(pk=job_id)
            locked_intake = StagedResourceIntake.objects.select_for_update().get(pk=intake_id)
            if locked_job.status != StagedResourceInspectionJobStatus.RUNNING:
                return False
            if (
                locked_intake.status != StagedResourceIntakeStatus.ANALYZING
                or self._has_intake_expired(locked_intake)
            ):
                locked_job.status = StagedResourceInspectionJobStatus.FAILED
                locked_job.error_code = "INTAKE_EXPIRED"
                locked_job.completed_at = timezone.now()
                locked_job.save(
                    update_fields=["status", "error_code", "completed_at", "updated_at"]
                )
                if self._has_intake_expired(locked_intake):
                    locked_intake.status = StagedResourceIntakeStatus.EXPIRED
                    locked_intake.cleanup_required = True
                    locked_intake.failure_code = "INTAKE_EXPIRED"
                    locked_intake.save(
                        update_fields=["status", "cleanup_required", "failure_code", "updated_at"]
                    )
                return False
            locked_job.status = StagedResourceInspectionJobStatus.COMPLETED
            locked_job.completed_at = timezone.now()
            locked_job.error_code = None
            locked_job.save(update_fields=["status", "completed_at", "error_code", "updated_at"])
            locked_intake.status = StagedResourceIntakeStatus.READY_FOR_PUBLICATION
            locked_intake.inspection_json = inspection
            locked_intake.verified_checksum = str(inspection["content_hash"])
            locked_intake.checksum_verification_status = "VERIFIED"
            locked_intake.failure_code = None
            locked_intake.save(
                update_fields=[
                    "status",
                    "inspection_json",
                    "verified_checksum",
                    "checksum_verification_status",
                    "failure_code",
                    "updated_at",
                ]
            )

            register_intake_object(locked_intake)
        return True

    @staticmethod
    def _recover_abandoned_jobs(*, intake_id: uuid.UUID | None = None) -> int:
        """Requeue inspections whose worker lease elapsed before terminal persistence.

        Args:
            intake_id: Optional exact intake; leave unrelated leases untouched.

        Returns:
            Number of jobs returned to the queue.
        """
        lease_cutoff = timezone.now() - timedelta(
            seconds=max(
                int(getattr(settings, "RESOURCE_REGISTRY_INTAKE_INSPECTION_LEASE_SECONDS", 900)), 60
            )
        )
        candidates = StagedResourceInspectionJob.objects.all()
        if intake_id is not None:
            candidates = candidates.filter(intake_id=intake_id)
        with transaction.atomic():
            jobs = tuple(
                candidates.select_for_update(skip_locked=True)
                .filter(
                    status=StagedResourceInspectionJobStatus.RUNNING,
                    started_at__lte=lease_cutoff,
                )
                .select_related("intake")
            )
            for job in jobs:
                job.status = StagedResourceInspectionJobStatus.QUEUED
                job.error_code = "INTAKE_INSPECTION_LEASE_EXPIRED"
                job.started_at = None
                job.save(update_fields=["status", "error_code", "started_at", "updated_at"])
                if job.intake.status == StagedResourceIntakeStatus.ANALYZING:
                    job.intake.status = StagedResourceIntakeStatus.COMPLETED
                    job.intake.save(update_fields=["status", "updated_at"])
        return len(jobs)

    @staticmethod
    def _has_intake_expired(intake: StagedResourceIntake) -> bool:
        """Return whether an active staging intake passed its retention deadline."""
        return intake.expires_at is not None and intake.expires_at <= timezone.now()

    @staticmethod
    def _safe_sheet_summary(*, name: str, context: dict[str, Any]) -> dict[str, object]:
        """Build a sanitized inspection summary for one logical sheet.

        Args:
            name: Display name of the inspected sheet.
            context: File context produced by the context builder.

        Returns:
            Minimal non-row-level metadata about the sheet.
        """
        summary = context.get("summary", {})
        schema = context.get("schema", [])
        return {
            "name": name,
            "status": "READY",
            "row_count": int(summary.get("row_count", 0)),
            "column_count": int(summary.get("column_count", 0)),
            "columns": [
                str(column.get("name", "")) for column in schema if isinstance(column, dict)
            ],
        }

    @staticmethod
    def calculate_sha256_file(path: Path) -> str:
        """Calculate a SHA-256 digest without loading the file into memory.

        Args:
            path: Local path to the downloaded staged object.

        Returns:
            Lower-case hexadecimal SHA-256 digest.
        """
        digest = hashlib.sha256()
        with path.open("rb") as source:
            for block in iter(lambda: source.read(1024 * 1024), b""):
                digest.update(block)
        return digest.hexdigest()
