"""Verify external-source formats, workbook transactions, and authenticated API contracts."""

from pathlib import Path
from typing import cast
from uuid import uuid4

import pyarrow as pa  # type: ignore[import-untyped]
import pyarrow.parquet as pq  # type: ignore[import-untyped]
import pytest
from django.db import connection
from django.template.loader import render_to_string
from django.test.utils import CaptureQueriesContext
from django.urls import resolve
from guardian.shortcuts import remove_perm
from openpyxl import Workbook  # type: ignore[import-untyped]
from rest_framework.test import APIRequestFactory, force_authenticate

from resource_registry.api.management.serializers import DatasetChangeResponseSerializer
from resource_registry.api.management.views import DatasetExternalSourceDetailView
from resource_registry.management_domain.changes import DatasetChangeService
from resource_registry.management_domain.errors import DatasetManagementError
from resource_registry.management_domain.external_sources import (
    DatasetExternalSourceService,
    source_payload,
)
from resource_registry.management_domain.worker import DatasetChangeWorker
from resource_registry.models import Dataset, DatasetChange
from resource_registry.tests.test_dataset_changes import create_verified_intake
from resource_registry.tests.test_dataset_maintenance_flows import Scope
from resource_registry.tests.test_dataset_management import (  # noqa: F401
    ManagementRequest,
    management_scope,
)
from resource_registry.tests.test_external_source_contract import (  # noqa: F401
    prepare_source,
    publish_source,
    source_scope,
)


def prepare_path(
    scope: Scope,
    path: Path,
    specification: dict[str, object],
    *,
    version: str = "v1",
    source_revision: int = 0,
) -> DatasetChange:
    """Prepare a binary or textual fixture using the verified intake contract."""
    _, editor, organization, namespace, _ = scope
    intake = create_verified_intake(editor, path)
    if path.suffix == ".ndjson":
        intake.resolved_file_format = "JSONL"
        intake.save(update_fields=["resolved_file_format"])
    change = DatasetChangeService(editor, organization).create(
        namespace_id=namespace.id,
        intake_id=intake.id,
        idempotency_key=str(uuid4()),
        specification=specification,
        source_reference={
            "provider": "sharepoint",
            "external_id": "drive:formats",
            "version": version,
        },
        expected_source_revision=source_revision,
    )
    DatasetChangeWorker().run_once(change_id=change.id)
    change.refresh_from_db()
    return change


@pytest.mark.parametrize("extension", ["csv", "xlsx", "json", "jsonl", "ndjson", "parquet"])
def test_supported_formats_publish_real_canonical_outputs(
    source_scope: Scope, tmp_path: Path, extension: str
) -> None:
    """Ingest supported file contents rather than a manifest describing them."""
    path = tmp_path / f"sales.{extension}"
    specification: dict[str, object] = {"name": "Sales"}
    if extension == "xlsx":
        book = Workbook()
        sheet = book.active
        sheet.title = "Sales"
        sheet.append(["id", "amount"])
        sheet.append([1, 10])
        book.save(path)
        book.close()
        specification["sheets"] = [{"sheet_name": "Sales", "selected": True}]
    elif extension == "parquet":
        pq.write_table(pa.table({"id": [1], "amount": [10]}), path)
    elif extension == "json":
        path.write_text('[{"id":1,"amount":10}]')
    elif extension in {"jsonl", "ndjson"}:
        path.write_text('{"id":1,"amount":10}\n')
    else:
        path.write_text("id,amount\n1,10\n")
    change = prepare_path(source_scope, path, specification)
    assert change.status == "READY", change.error_message
    dataset = publish_source(source_scope, change)
    assert dataset.current_version is not None
    assert dataset.current_version.profiles.get().metrics_json["row_count"] == 1
    assert dataset.external_source_id == change.external_source_id
    change.refresh_from_db()
    assert change.result["items"][0]["outcome"] == "created"
    if extension == "xlsx":
        assert change.result["items"][0]["sheet_name"] == "Sales"
        assert change.result["workbook_id"] == str(dataset.workbook_id)


def workbook_file(path: Path, sheets: list[str]) -> Path:
    """Build a small workbook with deterministic data for every requested sheet."""
    book = Workbook()
    book.remove(book.active)
    for name in sheets:
        sheet = book.create_sheet(name)
        sheet.append(["id", "amount"])
        sheet.append([1, 10])
    book.save(path)
    book.close()
    return path


def test_workbook_mapping_addition_removal_and_source_deactivation(
    source_scope: Scope, tmp_path: Path
) -> None:
    """Approve a whole workbook including a rename, new sheet, and removal atomically."""
    manager, editor, organization, _, _ = source_scope
    first = prepare_path(
        source_scope,
        workbook_file(tmp_path / "first.xlsx", ["Sales", "Costs"]),
        {"name": "Book", "sheets": [{"sheet_name": "Sales"}, {"sheet_name": "Costs"}]},
    )
    publish_source(source_scope, first)
    first.refresh_from_db()
    source = first.external_source
    assert source is not None
    source.refresh_from_db()
    assert source.workbook is not None
    sales = source.datasets.get(workbook_sheet="Sales")
    costs = source.datasets.get(workbook_sheet="Costs")
    change = prepare_path(
        source_scope,
        workbook_file(tmp_path / "next.xlsx", ["Revenue", "Regions"]),
        {
            "workbook_id": str(source.workbook_id),
            "workbook_revision": source.workbook.revision,
            "sheets": [
                {
                    "sheet_name": "Revenue",
                    "dataset_id": str(sales.id),
                    "expected_revision": sales.revision,
                },
                {"sheet_name": "Regions"},
            ],
            "archive": [{"dataset_id": str(costs.id), "expected_revision": costs.revision}],
        },
        version="v2",
        source_revision=source.revision,
    )
    assert change.status == "NEEDS_MANAGER", change.error_message
    with pytest.raises(DatasetManagementError):
        DatasetChangeService(editor, organization).publish(
            change.id, expected_revision=change.revision
        )
    result = DatasetChangeService(manager, organization).publish(
        change.id, expected_revision=change.revision
    )
    sales.refresh_from_db()
    costs.refresh_from_db()
    assert sales.workbook_sheet == "Revenue"
    assert costs.status == "ARCHIVED"
    assert len(result.result["items"]) == 3
    source.refresh_from_db()
    DatasetExternalSourceService(editor, organization).deactivate(
        source.id,
        expected_revision=source.revision,
        reason="Book removed",
        idempotency_key="book-removed",
    )
    assert not source.datasets.exclude(status="ARCHIVED").exists()


@pytest.mark.parametrize("prefix", ["api/v1", "auth/api/v1"])
def test_api_routes_expose_control_metadata_and_idempotent_deactivation(
    source_scope: Scope, tmp_path: Path, prefix: str
) -> None:
    """Exercise both mounted route trees with a nonstaff integration identity."""
    _, editor, organization, _, _ = source_scope
    editor.is_staff = False
    editor.save(update_fields=["is_staff"])
    first = prepare_source(source_scope, tmp_path / "first.csv")
    publish_source(source_scope, first)
    source = first.external_source
    assert source is not None
    source.refresh_from_db()
    path = f"/{prefix}/resource-registry/dataset-sources/{source.id}/"
    request = APIRequestFactory().get(path)
    cast(ManagementRequest, request).organization = organization
    force_authenticate(request, user=editor)
    route = resolve(path)
    result = route.func(request, **route.kwargs)
    assert result.status_code == 200
    assert result.data["last_applied_version"] == "v1"
    body = {"expected_revision": source.revision, "reason": "Origin gone"}
    for _ in range(2):
        request = APIRequestFactory().post(
            path + "deactivate/", body, format="json", HTTP_IDEMPOTENCY_KEY="api-gone"
        )
        cast(ManagementRequest, request).organization = organization
        force_authenticate(request, user=editor)
        route = resolve(path + "deactivate/")
        response = route.func(request, **route.kwargs)
        assert response.status_code == 200
        assert response.data["status"] == "INACTIVE"


def test_revoked_sync_permission_blocks_prepared_publication(
    source_scope: Scope, tmp_path: Path
) -> None:
    """Recheck source authority at publication, not only proposal admission."""
    _, editor, organization, namespace, _ = source_scope
    change = prepare_source(source_scope, tmp_path / "first.csv")
    remove_perm("resource_registry.synchronize_dataset_sources", editor, namespace)
    with pytest.raises(DatasetManagementError) as failure:
        DatasetChangeService(editor, organization).publish(
            change.id, expected_revision=change.revision
        )
    assert failure.value.code == "source_sync_denied"
    assert not Dataset.objects.filter(external_source_id=change.external_source_id).exists()


def test_foreign_source_is_hidden_from_api(source_scope: Scope, tmp_path: Path) -> None:
    """Reject a source lookup without the required namespace synchronization grant."""
    _, editor, organization, namespace, _ = source_scope
    change = prepare_source(source_scope, tmp_path / "first.csv")
    remove_perm("resource_registry.synchronize_dataset_sources", editor, namespace)
    request = APIRequestFactory().get("/source/")
    cast(ManagementRequest, request).organization = organization
    force_authenticate(request, user=editor)
    response = DatasetExternalSourceDetailView.as_view()(
        request, source_id=change.external_source_id
    )
    assert response.status_code == 403
    assert "external_id" not in response.data


def test_control_read_query_count_does_not_scale_per_sheet(
    source_scope: Scope, tmp_path: Path
) -> None:
    """Serialize source output metadata with a bounded number of database requests."""
    small_change = prepare_source(source_scope, tmp_path / "small.csv")
    publish_source(source_scope, small_change)
    small_source = small_change.external_source
    assert small_source is not None
    with CaptureQueriesContext(connection) as small_queries:
        small_result = source_payload(small_source)
    assert len(small_result["items"]) == 1
    change = prepare_path(
        source_scope,
        workbook_file(tmp_path / "many.xlsx", [f"Sheet{i}" for i in range(8)]),
        {"name": "Book", "sheets": [{"sheet_name": f"Sheet{i}"} for i in range(8)]},
    )
    publish_source(source_scope, change)
    source = change.external_source
    assert source is not None
    source.refresh_from_db()
    with CaptureQueriesContext(connection) as queries:
        result = source_payload(source)
    assert len(result["items"]) == 8
    assert len(queries) <= len(small_queries) + 1


def test_admin_detail_does_not_offer_manual_content_replacement(
    source_scope: Scope, tmp_path: Path
) -> None:
    """Explain source ownership while retaining metadata controls."""
    dataset = publish_source(source_scope, prepare_source(source_scope, tmp_path / "first.csv"))
    html = render_to_string(
        "admin/resource_registry/dataset/detail.html", {"dataset": dataset, "editable": True}
    )
    assert "Content is managed by an external source" in html
    assert "Update data" not in html
    assert "Save metadata" in html


@pytest.mark.parametrize("prefix", ["api/v1", "auth/api/v1"])
def test_create_and_publish_source_through_rest_returns_valid_contract(
    source_scope: Scope, tmp_path: Path, prefix: str
) -> None:
    """Connect request serialization, idempotent creation, worker output and response schemas."""

    _, editor, organization, namespace, _ = source_scope
    path = tmp_path / "api.csv"
    path.write_text("id,amount\n1,10\n")
    intake = create_verified_intake(editor, path)
    payload = {
        "namespace_id": str(namespace.id),
        "intake_id": str(intake.id),
        "specification": {"name": "Sales"},
        "source_reference": {"provider": "sharepoint", "external_id": "drive:api", "version": "v1"},
        "expected_source_revision": 0,
    }
    url = f"/{prefix}/resource-registry/dataset-changes/"
    identifiers = []
    for _ in range(2):
        request = APIRequestFactory().post(
            url, payload, format="json", HTTP_IDEMPOTENCY_KEY="create-api"
        )
        cast(ManagementRequest, request).organization = organization
        force_authenticate(request, user=editor)
        route = resolve(url)
        response = route.func(request, **route.kwargs)
        assert response.status_code == 201
        schema = DatasetChangeResponseSerializer(data=response.data)
        assert schema.is_valid(), schema.errors
        identifiers.append(response.data["id"])
    assert identifiers[0] == identifiers[1]
    change = DatasetChange.objects.get(pk=identifiers[0])
    DatasetChangeWorker().run_once(change_id=change.id)
    change.refresh_from_db()
    publish_url = f"{url}{change.id}/publish/"
    request = APIRequestFactory().post(
        publish_url, {"expected_revision": change.revision}, format="json"
    )
    cast(ManagementRequest, request).organization = organization
    force_authenticate(request, user=editor)
    route = resolve(publish_url)
    response = route.func(request, **route.kwargs)
    assert response.status_code == 200
    schema = DatasetChangeResponseSerializer(data=response.data)
    assert schema.is_valid(), schema.errors
    assert response.data["result"]["items"][0]["outcome"] == "created"
