from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook

from resource_registry.services.excel_dataset_adapter import (
    EmptyExcelWorksheetError,
    ExcelColumnLimitExceededError,
    ExcelDatasetAdapter,
    ExcelFileSizeLimitExceededError,
    ExcelRowLimitExceededError,
    UnsupportedExcelWorkbookError,
    WorksheetNotFoundError,
)

EXPECTED_ROW_COUNT = 2
EXPECTED_COLUMN_COUNT = 3


def _create_xlsx(path: Path, *, sheet_name: str = "Sheet1") -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_name
    worksheet.append(["customer_id", "country", "total_spend"])
    worksheet.append([1, "Costa Rica", 100.5])
    worksheet.append([2, "Mexico", 200.0])
    workbook.save(path)
    workbook.close()


def test_normalize_xlsx_to_csv(tmp_path: Path) -> None:
    source_path = tmp_path / "sample.xlsx"
    _create_xlsx(source_path)
    adapter = ExcelDatasetAdapter()

    with adapter.materialize_to_csv(source_path) as normalized:
        normalized_path = Path(normalized.normalized_path)
        assert normalized_path.exists()
        assert normalized.sheet_name == "Sheet1"
        assert normalized.row_count == EXPECTED_ROW_COUNT
        assert normalized.column_count == EXPECTED_COLUMN_COUNT
        assert normalized.columns == ("customer_id", "country", "total_spend")
        assert normalized.detected_header_row_index == 0
        assert normalized.header_row_span == 1
        assert normalized.skipped_rows == 0
        text = normalized_path.read_text(encoding="utf-8")
        assert "Costa Rica" in text
        assert "total_spend" in text

    assert normalized_path.exists() is False


def test_empty_worksheet_rejected(tmp_path: Path) -> None:
    source_path = tmp_path / "empty.xlsx"
    workbook = Workbook()
    workbook.save(source_path)
    workbook.close()
    adapter = ExcelDatasetAdapter()

    with pytest.raises(EmptyExcelWorksheetError):
        with adapter.materialize_to_csv(source_path):
            pass


def test_sheet_name_selection(tmp_path: Path) -> None:
    source_path = tmp_path / "multi.xlsx"
    workbook = Workbook()
    first = workbook.active
    first.title = "Summary"
    first.append(["id", "name"])
    first.append([1, "A"])
    detail = workbook.create_sheet("Detail")
    detail.append(["segment", "monto"])
    detail.append(["VIP", 300])
    workbook.save(source_path)
    workbook.close()

    adapter = ExcelDatasetAdapter()
    with adapter.materialize_to_csv(source_path, sheet_name="Detail") as normalized:
        assert normalized.sheet_name == "Detail"
        assert normalized.columns == ("segment", "monto")


def test_list_worksheets_returns_workbook_order(tmp_path: Path) -> None:
    source_path = tmp_path / "multi.xlsx"
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary.append(["id"])
    summary.append([1])
    workbook.create_sheet("Detail")
    workbook.create_sheet("Pivot")
    workbook.save(source_path)
    workbook.close()

    adapter = ExcelDatasetAdapter()
    assert adapter.list_worksheets(source_path) == ["Summary", "Detail", "Pivot"]


def test_missing_sheet_name_returns_worksheet_not_found(tmp_path: Path) -> None:
    source_path = tmp_path / "multi.xlsx"
    _create_xlsx(source_path, sheet_name="Summary")
    adapter = ExcelDatasetAdapter()

    with pytest.raises(WorksheetNotFoundError):
        adapter.normalize(source_path, sheet_name="Missing")


def test_duplicate_headers_are_deduplicated(tmp_path: Path) -> None:
    source_path = tmp_path / "dup.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["name", "name", "name"])
    worksheet.append(["A", "B", "C"])
    workbook.save(source_path)
    workbook.close()
    adapter = ExcelDatasetAdapter()

    with adapter.materialize_to_csv(source_path) as normalized:
        assert normalized.columns == ("name", "name_2", "name_3")


def test_empty_headers_get_fallback_names(tmp_path: Path) -> None:
    source_path = tmp_path / "fallback.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["", None, "value"])
    worksheet.append([1, 2, 3])
    workbook.save(source_path)
    workbook.close()
    adapter = ExcelDatasetAdapter()

    with adapter.materialize_to_csv(source_path) as normalized:
        assert normalized.columns == ("column_1", "column_2", "value")
        assert normalized.fallback_name_reasons["column_1"] == "empty_header_cell"
        assert normalized.fallback_name_reasons["column_2"] == "empty_header_cell"


def test_row_limit_enforced(tmp_path: Path) -> None:
    source_path = tmp_path / "rows.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["id"])
    worksheet.append([1])
    worksheet.append([2])
    workbook.save(source_path)
    workbook.close()
    adapter = ExcelDatasetAdapter(max_rows=2)

    with pytest.raises(ExcelRowLimitExceededError):
        adapter.normalize(source_path)


def test_column_limit_enforced(tmp_path: Path) -> None:
    source_path = tmp_path / "cols.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["a", "b", "c"])
    worksheet.append([1, 2, 3])
    workbook.save(source_path)
    workbook.close()
    adapter = ExcelDatasetAdapter(max_columns=2)

    with pytest.raises(ExcelColumnLimitExceededError):
        adapter.normalize(source_path)


def test_formula_cells_use_cached_values_only(tmp_path: Path) -> None:
    source_path = tmp_path / "formula.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["value", "formula"])
    worksheet.append([10, "=A2*2"])
    workbook.save(source_path)
    workbook.close()
    adapter = ExcelDatasetAdapter()

    with adapter.materialize_to_csv(source_path) as normalized:
        serialized = Path(normalized.normalized_path).read_text(encoding="utf-8")
        assert "=A2*2" not in serialized


def test_file_size_limit_enforced(tmp_path: Path) -> None:
    source_path = tmp_path / "large.xlsx"
    _create_xlsx(source_path)
    adapter = ExcelDatasetAdapter(max_file_bytes=1)

    with pytest.raises(ExcelFileSizeLimitExceededError):
        adapter.normalize(source_path)


def test_encrypted_workbook_error_mapping(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    source_path = tmp_path / "encrypted.xlsx"
    source_path.write_text("not-an-excel", encoding="utf-8")
    adapter = ExcelDatasetAdapter()

    def _raise(*args, **kwargs):
        _ = args, kwargs
        raise ValueError("Workbook is encrypted")

    monkeypatch.setattr(
        "resource_registry.services.excel_dataset_adapter.load_workbook",
        _raise,
    )
    with pytest.raises(UnsupportedExcelWorkbookError):
        adapter.normalize(source_path)


def test_header_detection_skips_title_rows(tmp_path: Path) -> None:
    source_path = tmp_path / "title_rows.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["Campaign Segment Summary"])
    worksheet.append(["Generated on", "2026-07-08"])
    worksheet.append(["customer_id", "country", "total_spend"])
    worksheet.append([1, "Costa Rica", 100.0])
    worksheet.append([2, "Mexico", 250.0])
    workbook.save(source_path)
    workbook.close()

    adapter = ExcelDatasetAdapter()
    with adapter.materialize_to_csv(source_path) as normalized:
        assert normalized.detected_header_row_index == 2
        assert normalized.columns == ("customer_id", "country", "total_spend")
        assert normalized.row_count == 2


def test_grouped_headers_are_flattened(tmp_path: Path) -> None:
    source_path = tmp_path / "grouped_headers.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["Segment", "Segment", "Metrics", "Metrics"])
    worksheet.append(["name", "region", "score", "confidence"])
    worksheet.append(["GM Owners", "US", 0.88, 0.67])
    workbook.save(source_path)
    workbook.close()

    adapter = ExcelDatasetAdapter()
    with adapter.materialize_to_csv(source_path) as normalized:
        assert normalized.header_row_span == 2
        assert normalized.columns == (
            "segment_name",
            "segment_region",
            "metrics_score",
            "metrics_confidence",
        )
        assert normalized.row_count == 1


def test_segment_label_rows_not_used_as_headers(tmp_path: Path) -> None:
    source_path = tmp_path / "segment_labels.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(
        [
            "GM Owners",
            "GM Owners with OnStar Propensity",
            "Cross Shoppers with Propensities",
        ],
    )
    worksheet.append(["generated", "for", "analysis"])
    worksheet.append(["customer_id", "segment", "propensity_score"])
    worksheet.append([1, "GM Owners", 0.91])
    worksheet.append([2, "Cross Shoppers", 0.37])
    workbook.save(source_path)
    workbook.close()

    adapter = ExcelDatasetAdapter()
    with adapter.materialize_to_csv(source_path) as normalized:
        assert normalized.detected_header_row_index == 2
        assert normalized.columns == ("customer_id", "segment", "propensity_score")
        assert normalized.columns[0] != "gm_owners"


def test_fully_empty_columns_are_pruned_from_schema(tmp_path: Path) -> None:
    source_path = tmp_path / "empty_columns.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["customer_id", None, "", "country", ""])
    worksheet.append([1, None, "", "Costa Rica", ""])
    worksheet.append([2, None, "", "Mexico", ""])
    workbook.save(source_path)
    workbook.close()

    adapter = ExcelDatasetAdapter()
    with adapter.materialize_to_csv(source_path) as normalized:
        assert normalized.columns == ("customer_id", "country")
        assert normalized.column_count == 2
        assert any("Dropped" in warning for warning in normalized.detection_warnings)
