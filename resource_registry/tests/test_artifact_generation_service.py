from __future__ import annotations

import io
import uuid
from dataclasses import dataclass

from django.contrib.auth import get_user_model
from django.test import TestCase
from openpyxl import load_workbook

from resource_registry.models import (
    Dataset,
    DatasetExportPlan,
    DatasetGeneratedArtifact,
    DatasetSourceType,
    DatasetVersion,
    DatasetVersionStatus,
)
from resource_registry.read_auth.context import ReadAuthorizationContext
from resource_registry.read_auth.exceptions import DataRegistryAuthorizationError
from resource_registry.services.artifact_generation_service import (
    ArtifactGenerationService,
    CSVArtifactGenerator,
    XlsxArtifactGenerator,
)
from resource_registry.storage.base import StorageObjectMetadata

User = get_user_model()


@dataclass
class _FakeStorageProvider:
    calls: list[tuple[str, str]]
    payloads: list[bytes]

    def exists(self, uri_or_path: str) -> bool:
        _ = uri_or_path
        return False

    def head(self, uri_or_path: str):
        _ = uri_or_path
        raise NotImplementedError

    def put(self, local_path: str, destination_path: str) -> StorageObjectMetadata:
        self.calls.append((local_path, destination_path))
        with open(local_path, "rb") as generated:
            self.payloads.append(generated.read())
        return StorageObjectMetadata(
            uri=f"local://{destination_path}",
            provider="local",
            size_bytes=1234,
            content_type="application/octet-stream",
            checksum="checksum-1",
            metadata={"safe": "ok"},
        )

    def get_uri(self, destination_path: str) -> str:
        return f"local://{destination_path}"

    def delete(self, uri_or_path: str) -> None:
        _ = uri_or_path

    def download_to_file(self, uri_or_path: str, local_path: str):
        _ = uri_or_path, local_path
        raise NotImplementedError


class _FakeStorageFactory:
    def __init__(self, provider: _FakeStorageProvider) -> None:
        self.provider = provider
        self.calls: list[str | None] = []

    def create_provider(self, provider_key=None):
        self.calls.append(provider_key)
        return self.provider


class _FakeQueryExecutor:
    def __init__(self, payload: dict | None = None) -> None:
        self.payload = payload or {
            "columns": ["country", "total_spend"],
            "rows": [["CR", 100], ["US", 80]],
            "row_count": 2,
            "truncated": False,
        }
        self.calls: list[dict] = []

    def execute_query(self, **kwargs):
        self.calls.append(kwargs)
        return self.payload


class ArtifactGenerationServiceTests(TestCase):
    def setUp(self) -> None:
        self.user = User.objects.create_user(username="generator_user", password="unused")
        self.tenant_id = uuid.uuid4()
        self.dataset = Dataset.objects.create(
            tenant_id=self.tenant_id,
            slug="sales-dataset",
            name="Sales Dataset",
            source_type=DatasetSourceType.FILE_UPLOAD,
            created_by=self.user,
        )
        self.version = DatasetVersion.objects.create(
            dataset=self.dataset,
            version_number=1,
            source_uri="local://sales.csv",
            content_hash="hash-1",
            status=DatasetVersionStatus.READY,
        )
        self.context = ReadAuthorizationContext(tenant_id=str(self.tenant_id), roles=("admin",))

    def _export_plan(self, artifact_type: str, filename: str) -> DatasetExportPlan:
        return DatasetExportPlan.objects.create(
            tenant_id=self.tenant_id,
            dataset=self.dataset,
            version=self.version,
            artifact_type=artifact_type,
            title="Dataset Export",
            export_plan_json={
                "filename": filename,
                "data_requirements": [
                    {
                        "name": "spend_by_country",
                        "sql": (
                            "SELECT country, SUM(total_spend) AS total_spend "
                            "FROM dataset GROUP BY country ORDER BY total_spend DESC LIMIT 20"
                        ),
                        "columns": ["country", "total_spend"],
                        "max_rows": 20,
                    },
                ],
                "sheets": [{"name": "Summary", "type": "table", "source": "spend_by_country"}],
                "sections": [{"id": "section_1", "title": "Summary", "source": "spend_by_country"}],
                "slides": [{"id": "slide_1", "title": "Summary", "source": "spend_by_country"}],
            },
            status="SUCCEEDED",
            created_by=self.user,
        )

    def _service(
        self,
    ) -> tuple[ArtifactGenerationService, _FakeStorageProvider, _FakeQueryExecutor]:
        provider = _FakeStorageProvider(calls=[], payloads=[])
        factory = _FakeStorageFactory(provider)
        executor = _FakeQueryExecutor()
        service = ArtifactGenerationService(
            query_executor=executor,
            storage_factory=factory,
        )
        return service, provider, executor

    def test_csv_generation(self) -> None:
        plan = self._export_plan("csv", "campaign_report.csv")
        service, provider, executor = self._service()

        result = service.generate_from_export_plan(
            context=self.context,
            export_plan_id=str(plan.id),
            created_by=self.user,
        )

        self.assertTrue(result["success"])
        self.assertEqual(result["artifact"]["artifact_type"], "csv")
        self.assertEqual(len(provider.calls), 1)
        self.assertEqual(len(executor.calls), 1)

    def test_excel_generation(self) -> None:
        plan = self._export_plan("xlsx", "campaign_report.xlsx")
        service, provider, _ = self._service()
        result = service.generate_from_export_plan(
            context=self.context,
            export_plan_id=str(plan.id),
            created_by=self.user,
        )
        self.assertTrue(result["success"])
        self.assertEqual(result["artifact"]["artifact_type"], "xlsx")
        self.assertEqual(len(provider.calls), 1)
        workbook = load_workbook(io.BytesIO(provider.payloads[0]))
        self.assertIn("Summary", workbook.sheetnames)
        worksheet = workbook["Summary"]
        self.assertEqual(worksheet.freeze_panes, "A2")
        self.assertIsNotNone(worksheet.auto_filter.ref)
        self.assertTrue(worksheet["A1"].font.bold)
        self.assertEqual(worksheet["A1"].value, "country")
        self.assertEqual(worksheet["A2"].value, "CR")
        self.assertEqual(worksheet["B2"].value, 100)
        self.assertEqual(
            result["artifact"]["mime_type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    def test_pdf_generation(self) -> None:
        plan = self._export_plan("pdf", "campaign_report.pdf")
        service, provider, _ = self._service()
        result = service.generate_from_export_plan(
            context=self.context,
            export_plan_id=str(plan.id),
            created_by=self.user,
        )
        self.assertFalse(result["success"])
        self.assertEqual(result["code"], "unsupported_artifact_type")
        self.assertEqual(len(provider.calls), 0)

    def test_pptx_generation(self) -> None:
        plan = self._export_plan("pptx", "campaign_report.pptx")
        service, provider, _ = self._service()
        result = service.generate_from_export_plan(
            context=self.context,
            export_plan_id=str(plan.id),
            created_by=self.user,
        )
        self.assertFalse(result["success"])
        self.assertEqual(result["code"], "unsupported_artifact_type")
        self.assertEqual(len(provider.calls), 0)

    def test_storage_persistence_creates_dataset_generated_artifact(self) -> None:
        plan = self._export_plan("csv", "campaign_report.csv")
        service, _, _ = self._service()
        _ = service.generate_from_export_plan(
            context=self.context,
            export_plan_id=str(plan.id),
            created_by=self.user,
        )
        self.assertEqual(DatasetGeneratedArtifact.objects.count(), 1)

    def test_invalid_export_plan_not_found(self) -> None:
        service, _, _ = self._service()
        with self.assertRaises(DataRegistryAuthorizationError):
            service.generate_from_export_plan(
                context=self.context,
                export_plan_id=str(uuid.uuid4()),
                created_by=self.user,
            )

    def test_unsupported_artifact_type(self) -> None:
        plan = self._export_plan("xml", "campaign_report.xml")
        service, provider, executor = self._service()
        result = service.generate_from_export_plan(
            context=self.context,
            export_plan_id=str(plan.id),
            created_by=self.user,
        )
        self.assertFalse(result["success"])
        self.assertEqual(result["code"], "unsupported_artifact_type")
        self.assertEqual(len(provider.calls), 0)
        self.assertEqual(len(executor.calls), 0)

    def test_strategy_pattern_generator_registry(self) -> None:
        service, _, _ = self._service()
        self.assertIsInstance(service._generators["csv"], CSVArtifactGenerator)
        self.assertIn("xlsx", service._generators)
        self.assertNotIn("pdf", service._generators)
        self.assertNotIn("pptx", service._generators)

    def test_generated_metadata_shape_and_no_filesystem_leakage(self) -> None:
        plan = self._export_plan("csv", "campaign_report.csv")
        service, _, _ = self._service()
        result = service.generate_from_export_plan(
            context=self.context,
            export_plan_id=str(plan.id),
            created_by=self.user,
        )
        artifact = result["artifact"]
        self.assertIn("artifact_id", artifact)
        self.assertIn("storage_uri", artifact)
        self.assertFalse(str(artifact["storage_uri"]).startswith("/"))
        serialized = str(artifact).lower()
        self.assertNotIn("token", serialized)
        self.assertNotIn("password", serialized)
        self.assertNotIn("secret", serialized)
        self.assertNotIn("sig=", serialized)

    def test_xlsx_formula_injection_values_are_escaped(self) -> None:
        plan = self._export_plan("xlsx", "campaign_report.xlsx")
        provider = _FakeStorageProvider(calls=[], payloads=[])
        factory = _FakeStorageFactory(provider)
        executor = _FakeQueryExecutor(
            payload={
                "columns": ["country", "notes"],
                "rows": [["CR", "=SUM(1,1)"], ["US", "@danger"], ["DE", "-10"]],
                "row_count": 3,
                "truncated": False,
            },
        )
        service = ArtifactGenerationService(query_executor=executor, storage_factory=factory)

        result = service.generate_from_export_plan(
            context=self.context,
            export_plan_id=str(plan.id),
            created_by=self.user,
        )

        self.assertTrue(result["success"])
        workbook = load_workbook(io.BytesIO(provider.payloads[0]))
        worksheet = workbook["Summary"]
        self.assertEqual(worksheet["B2"].value, "'=SUM(1,1)")
        self.assertEqual(worksheet["B3"].value, "'@danger")
        self.assertEqual(worksheet["B4"].value, "'-10")

    def test_xlsx_artifact_too_large_when_sheet_limits_exceeded(self) -> None:
        plan = self._export_plan("xlsx", "campaign_report.xlsx")
        provider = _FakeStorageProvider(calls=[], payloads=[])
        factory = _FakeStorageFactory(provider)
        executor = _FakeQueryExecutor(
            payload={
                "columns": ["country", "total_spend"],
                "rows": [["CR", 100], ["US", 80]],
                "row_count": 2,
                "truncated": False,
            },
        )
        service = ArtifactGenerationService(
            query_executor=executor,
            storage_factory=factory,
            generators={
                "csv": CSVArtifactGenerator(),
                "xlsx": XlsxArtifactGenerator(max_rows_per_sheet=1),
            },
        )

        result = service.generate_from_export_plan(
            context=self.context,
            export_plan_id=str(plan.id),
            created_by=self.user,
        )

        self.assertFalse(result["success"])
        self.assertEqual(result["status"], 413)
        self.assertEqual(result["code"], "artifact_too_large")
        self.assertEqual(len(provider.calls), 0)

    def test_generation_denied_when_dataset_not_in_allowed_scope(self) -> None:
        plan = self._export_plan("csv", "campaign_report.csv")
        service, provider, executor = self._service()
        denied_context = ReadAuthorizationContext(
            tenant_id=str(self.tenant_id),
            roles=("reader",),
            allowed_dataset_ids=frozenset({str(uuid.uuid4())}),
            metadata={"requested_tenant_id": str(self.tenant_id)},
        )

        with self.assertRaises(DataRegistryAuthorizationError):
            service.generate_from_export_plan(
                context=denied_context,
                export_plan_id=str(plan.id),
                created_by=self.user,
            )
        self.assertEqual(len(provider.calls), 0)
        self.assertEqual(len(executor.calls), 0)
