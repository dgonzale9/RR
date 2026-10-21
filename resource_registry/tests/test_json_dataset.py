"""Contract tests for strict JSON dataset ingestion."""

from pathlib import Path

import duckdb
import pytest
from django.test import override_settings
from openpyxl import Workbook, load_workbook

from resource_registry.services.artifact_generation_service import (
    CSVArtifactGenerator,
    XlsxArtifactGenerator,
)
from resource_registry.services.canonical_dataset_builder import CanonicalDatasetBuilder
from resource_registry.services.json_dataset import JsonDatasetValidationError, scan_json_dataset
from resource_registry.services.schema_detection import SchemaDetectionStrategyResolver


def test_json_dataset_unions_late_columns_and_preserves_nested_values(tmp_path: Path) -> None:
    source = tmp_path / "dataset.json"
    source.write_text('[{"payload":{"customer":{"id":"á"}}},{"late":[1,"two"]}]', encoding="utf-8")

    detected = SchemaDetectionStrategyResolver().detect(source_path=str(source), source_format="json")
    assert detected.columns == [{"name": "payload", "type": "JSON"}, {"name": "late", "type": "JSON"}]
    with CanonicalDatasetBuilder().build(source_path=str(source), source_format="json") as artifact:
        with duckdb.connect(database=":memory:") as connection:
            result = connection.execute(
                "SELECT json_extract_string(payload, '$.customer.id'), json_extract(late, '$[1]') FROM read_parquet(?) ORDER BY payload NULLS LAST",
                [artifact.canonical_path],
            ).fetchall()
    assert result == [("á", None), (None, '"two"')]


@pytest.mark.parametrize("content", ["{}", "[]", '{"one": 1}\n{"two": 2}', '[{"a": 1, "a": 2}]', '[{"a": NaN}]'])
def test_json_dataset_rejects_invalid_root_and_values(tmp_path: Path, content: str) -> None:
    source = tmp_path / "invalid.json"
    source.write_text(content, encoding="utf-8")
    with pytest.raises(JsonDatasetValidationError):
        scan_json_dataset(source)


@override_settings(RESOURCE_REGISTRY_JSON_MAX_RECORD_BYTES=32)
def test_json_dataset_rejects_an_oversized_record(tmp_path: Path) -> None:
    source = tmp_path / "large-record.json"
    source.write_text('[{"payload":"' + ("x" * 128) + '"}]', encoding="utf-8")
    with pytest.raises(JsonDatasetValidationError, match="JSON_RECORD_TOO_LARGE"):
        scan_json_dataset(source)


def test_json_values_export_as_json_and_reimport_semantically(tmp_path: Path) -> None:
    rows = {"dataset": [{"payload": {"name": "á", "items": [1, "two"]}}]}
    plan = {"filename": "export.csv", "data_requirements": [{"name": "dataset"}]}
    exported_csv = CSVArtifactGenerator().generate(export_plan=plan, data_results=rows)
    csv_source = tmp_path / "export.csv"
    csv_source.write_bytes(exported_csv.content)
    assert SchemaDetectionStrategyResolver().detect(
        source_path=str(csv_source), source_format="csv"
    ).columns == [{"name": "payload", "type": "JSON"}]

    exported_xlsx = XlsxArtifactGenerator().generate(
        export_plan={**plan, "sheets": [{"name": "Data", "source": "dataset"}]}, data_results=rows
    )
    xlsx_source = tmp_path / "export.xlsx"
    xlsx_source.write_bytes(exported_xlsx.content)
    workbook = load_workbook(xlsx_source, read_only=True)
    assert next(workbook.active.iter_rows(min_row=2, values_only=True))[0] == '{"name":"á","items":[1,"two"]}'
    workbook.close()


def test_xlsx_json_cells_are_detected_and_materialized_as_logical_json(tmp_path: Path) -> None:
    """Preserve valid XLSX JSON cells through schema detection and canonical Parquet."""
    source = tmp_path / "records.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["identifier", "payload"])
    worksheet.append(["one", '{"customer":{"id":"á"}}'])
    worksheet.append(["two", '[1,"two"]'])
    workbook.save(source)
    workbook.close()

    detected = SchemaDetectionStrategyResolver().detect(source_path=str(source), source_format="xlsx")
    assert {column["name"]: column["type"] for column in detected.columns}["payload"] == "JSON"
    with CanonicalDatasetBuilder().build(source_path=str(source), source_format="xlsx") as artifact:
        assert artifact.metadata["logical_json_columns"] == ["payload"]
        with duckdb.connect(database=":memory:") as connection:
            result = connection.execute(
                "SELECT json_extract_string(payload, '$.customer.id') FROM read_parquet(?) "
                "ORDER BY identifier",
                [artifact.canonical_path],
            ).fetchall()
    assert result == [("á",), (None,)]
