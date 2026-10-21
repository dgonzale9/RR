from __future__ import annotations

import base64
import io
from types import SimpleNamespace
from unittest.mock import MagicMock

import pytest
from django.test import override_settings
from openpyxl import load_workbook

from resource_registry.read_auth.context import ReadAuthorizationContext
from resource_registry.services.tabular_transformation_execution_service import (
    JobSpecValidationError,
    TabularTransformationExecutionService,
)

_EXPECTED_RECORD_COUNT = 3


def _minimal_job_spec(scope_type: str, scope_id: str, *, version: str = "1.0") -> dict:
    return {
        "job_type": "tabular_transformation",
        "version": version,
        "inputs": {"scope": {"type": scope_type, "id": scope_id}},
        "pipeline": [
            {"stage": "normalize", "config": {"trim_whitespace": True}},
            {
                "stage": "aggregate",
                "config": {"groupings": [{"name": "by_country", "dimensions": ["country"]}]},
            },
            {"stage": "narrate", "config": {"mode": "llm_with_fallback"}},
        ],
        "outputs": {
            "logical_artifacts": [{"name": "primary_report", "type": "table_bundle"}],
            "render": [{"format": "json"}, {"format": "xlsx"}],
        },
        "constraints": {
            "no_fabrication": True,
            "evidence_required": True,
            "max_output_rows": 10,
        },
        "rulesets": {
            "cleaning_policy_id": "default-cleaning-v1",
            "classification_policy_id": "default-classification-v1",
            "scoring_policy_id": "default-scoring-v1",
        },
    }


def test_execute_dataset_scope_returns_logical_and_rendered_outputs() -> None:
    read_facade = MagicMock()
    read_facade.get_schema.return_value = SimpleNamespace(
        schema_json=[{"name": "country"}, {"name": "spend"}]
    )
    query_executor = MagicMock()
    query_executor.execute_query.return_value = {
        "columns": ["country", "spend"],
        "rows": [["CR", 100], ["US", 200]],
        "row_count": 2,
        "truncated": False,
    }
    service = TabularTransformationExecutionService(
        read_facade=read_facade,
        query_executor=query_executor,
    )

    payload = service.execute(
        context=ReadAuthorizationContext(tenant_id="tenant-1"),
        job_spec=_minimal_job_spec("dataset", "dataset-1"),
    )

    assert payload["success"] is True
    assert payload["code"] == "transformation_completed"
    assert isinstance(payload["logical_result"]["logical_artifacts"], list)
    assert payload["logical_result"]["metadata"] == {}
    render_formats = {item["format"] for item in payload["rendered_outputs"]}
    assert render_formats == {"json", "xlsx"}
    assert payload["engine_mode"] == "legacy"


def test_execute_workbook_group_scope_uses_workbook_context_rows() -> None:
    read_facade = MagicMock()
    read_facade.get_workbook_context.return_value = SimpleNamespace(
        sheets=[
            SimpleNamespace(
                dataset_id="d-1",
                slug="sheet-a",
                name="Sheet A",
                sheet_name="Sheet A",
                row_count=10,
                column_count=4,
            ),
        ]
    )
    service = TabularTransformationExecutionService(read_facade=read_facade)
    job_spec = _minimal_job_spec("workbook_group", "group-1")
    payload = service.execute(
        context=ReadAuthorizationContext(tenant_id="tenant-1"),
        job_spec=job_spec,
    )
    assert payload["success"] is True
    assert payload["logical_result"]["scope"]["type"] == "workbook_group"


def test_execute_rejects_invalid_schema() -> None:
    service = TabularTransformationExecutionService()
    with pytest.raises(JobSpecValidationError):
        service.execute(
            context=ReadAuthorizationContext(tenant_id="tenant-1"),
            job_spec={"job_type": "tabular_transformation"},
        )


def test_execute_rejects_unknown_ruleset() -> None:
    service = TabularTransformationExecutionService()
    spec = _minimal_job_spec("query_result", "inline-1")
    spec["inputs"]["query_result"] = {"rows": [{"country": "CR"}]}
    spec["rulesets"]["classification_policy_id"] = "missing-policy"
    with pytest.raises(JobSpecValidationError):
        service.execute(
            context=ReadAuthorizationContext(tenant_id="tenant-1"),
            job_spec=spec,
        )


def test_v10_rejects_score_stage() -> None:
    service = TabularTransformationExecutionService()
    spec = _minimal_job_spec("query_result", "inline-1", version="1.0")
    spec["inputs"]["query_result"] = {"rows": [{"country": "CR", "index": 150}]}
    spec["pipeline"].insert(1, {"stage": "score", "config": {"rules": []}})
    with pytest.raises(JobSpecValidationError):
        service.execute(context=ReadAuthorizationContext(tenant_id="tenant-1"), job_spec=spec)


def test_v10_rejects_metadata_stage() -> None:
    service = TabularTransformationExecutionService()
    spec = _minimal_job_spec("query_result", "inline-1", version="1.0")
    spec["inputs"]["query_result"] = {"rows": [{"country": "CR"}]}
    spec["pipeline"].insert(1, {"stage": "metadata", "config": {"extractors": []}})
    with pytest.raises(JobSpecValidationError):
        service.execute(context=ReadAuthorizationContext(tenant_id="tenant-1"), job_spec=spec)


@override_settings(RESOURCE_REGISTRY_TRANSFORMATION_ENGINE_MODE="rules_v1_1")
def test_rules_v11_mode_executes_score_rank_and_metadata() -> None:

    service = TabularTransformationExecutionService()
    spec = _minimal_job_spec("query_result", "inline-1", version="1.1")
    spec["inputs"]["query_result"] = {
        "rows": [
            {
                "attribute_path": "a/b",
                "attribute_name": "coffee buyers",
                "index": 150,
                "overlap": 0.2,
                "report_line": "Report Date: 2026-07-01",
                "persona": "Explorers",
            },
            {
                "attribute_path": "a/c",
                "attribute_name": "tea readers",
                "index": 100,
                "overlap": 0.9,
                "report_line": "Report Date: 2026-07-01",
                "persona": "Explorers",
            },
            {
                "attribute_path": "x/z",
                "attribute_name": "other",
                "index": 80,
                "overlap": 0.4,
                "report_line": "Report Date: 2026-07-01",
                "persona": "Explorers",
            },
        ]
    }
    spec["outputs"]["logical_artifacts"] = [
        {"name": "ranked_signals", "type": "table"},
        {"name": "by_country", "type": "table"},
    ]
    spec["outputs"]["render"] = [
        {
            "format": "xlsx",
            "options": {
                "sheet_plan": {
                    "sheets": [
                        {
                            "name": "Ranked",
                            "source_artifact": "ranked_signals",
                            "columns": ["attribute_name", "index", "_segment"],
                            "sort": {"by": "index", "order": "desc"},
                            "limit": 5,
                            "freeze_header": True,
                            "title": "Ranked Signals",
                        },
                        {
                            "name": "ByCountry",
                            "source_artifact": "by_country",
                            "columns": ["country", "count"],
                            "freeze_header": False,
                        },
                    ]
                }
            },
        }
    ]
    spec["pipeline"] = [
        {
            "stage": "classify",
            "config": {
                "target_field": "bucket",
                "confidence_field": "bucket_confidence",
                "fallback_bucket": "all_other",
                "passes": [
                    {
                        "name": "path",
                        "rules": [
                            {
                                "rule_id": "path_buy",
                                "label": "WHAT I BUY",
                                "predicate": {
                                    "operator": "contains",
                                    "field": "attribute_path",
                                    "value": "/b",
                                },
                                "confidence": 0.95,
                            }
                        ],
                    },
                    {
                        "name": "keyword",
                        "rules": [
                            {
                                "rule_id": "kw_read",
                                "label": "WHAT I LOOK AT",
                                "predicate": {
                                    "operator": "contains",
                                    "fields": ["attribute_name", "attribute_path"],
                                    "value": "read",
                                },
                                "confidence": 0.7,
                            }
                        ],
                    },
                ],
            },
        },
        {
            "stage": "score",
            "config": {
                "output_label_field": "signal_class",
                "default_label": "average",
                "rules": [
                    {
                        "rule_id": "distinction_hi",
                        "label": "distinction",
                        "when": {"operator": "gt", "field": "index", "value": 120},
                    },
                    {
                        "rule_id": "scale_neutral_high_overlap",
                        "label": "scale",
                        "when": {
                            "operator": "and",
                            "conditions": [
                                {
                                    "operator": "between",
                                    "field": "index",
                                    "min": 95,
                                    "max": 105,
                                },
                                {"operator": "gt", "field": "overlap", "value": 0.8},
                            ],
                        },
                    },
                ],
            },
        },
        {
            "stage": "rank",
            "config": {
                "artifact_name": "ranked_signals",
                "segments": [
                    {
                        "name": "distinction",
                        "predicate": {
                            "operator": "contains",
                            "field": "signal_class",
                            "value": "distinction",
                        },
                        "sort_by": "index",
                        "order": "desc",
                        "limit": 2,
                        "pad_from_segment": "average",
                    },
                    {
                        "name": "average",
                        "predicate": {
                            "operator": "contains",
                            "field": "signal_class",
                            "value": "average",
                        },
                        "sort_by": "index",
                        "order": "desc",
                        "limit": 2,
                    },
                ],
            },
        },
        {
            "stage": "aggregate",
            "config": {
                "groupings": [{"name": "by_country", "dimensions": ["country"]}],
            },
        },
        {
            "stage": "metadata",
            "config": {
                "extractors": [
                    {
                        "extractor_id": "persona_name",
                        "type": "column_value",
                        "target_key": "persona_name",
                        "options": {"column": "persona", "mode": "unique"},
                    },
                    {
                        "extractor_id": "report_date",
                        "type": "regex",
                        "target_key": "report_date",
                        "options": {
                            "column": "report_line",
                            "pattern": r"Report Date:\s*(\d{4}-\d{2}-\d{2})",
                            "group": 1,
                        },
                    },
                    {
                        "extractor_id": "record_count",
                        "type": "aggregate",
                        "target_key": "record_count",
                        "options": {"operation": "count"},
                    },
                    {
                        "extractor_id": "source",
                        "type": "constant",
                        "target_key": "source",
                        "options": {"value": "epsilon_like"},
                    },
                ]
            },
        },
        {"stage": "narrate", "config": {"mode": "deterministic"}},
    ]

    payload = service.execute(
        context=ReadAuthorizationContext(tenant_id="tenant-1"),
        job_spec=spec,
    )

    assert payload["success"] is True
    assert payload["engine_mode"] == "rules_v1_1"
    assert payload["logical_result"]["metadata"]["persona_name"] == "Explorers"
    assert payload["logical_result"]["metadata"]["report_date"] == "2026-07-01"
    assert payload["logical_result"]["metadata"]["record_count"] == _EXPECTED_RECORD_COUNT

    trace = payload["decision_trace"]
    assert any(item["stage"] == "metadata" and item["rules"] for item in trace)
    assert any(item["stage"] == "score" and item["rules"] for item in trace)

    xlsx_payload = next(item for item in payload["rendered_outputs"] if item["format"] == "xlsx")
    content = base64.b64decode(xlsx_payload["content_base64"])
    workbook = load_workbook(io.BytesIO(content))
    assert set(workbook.sheetnames) == {"Ranked", "ByCountry"}
    assert workbook["Ranked"]["A1"].value == "Ranked Signals"


@override_settings(RESOURCE_REGISTRY_TRANSFORMATION_ENGINE_MODE="rules_v1_1")
def test_metadata_stage_fails_fast_when_evidence_missing() -> None:

    service = TabularTransformationExecutionService()
    spec = _minimal_job_spec("query_result", "inline-1", version="1.1")
    spec["inputs"]["query_result"] = {"rows": [{"field": "no date"}]}
    spec["pipeline"] = [
        {
            "stage": "metadata",
            "config": {
                "extractors": [
                    {
                        "extractor_id": "missing_regex",
                        "type": "regex",
                        "target_key": "report_date",
                        "options": {
                            "column": "field",
                            "pattern": r"(\\d{4}-\\d{2}-\\d{2})",
                            "group": 1,
                        },
                    }
                ]
            },
        }
    ]

    with pytest.raises(JobSpecValidationError):
        service.execute(
            context=ReadAuthorizationContext(tenant_id="tenant-1"),
            job_spec=spec,
        )


def test_sheet_plan_semantics_reject_duplicate_sheet_names() -> None:
    service = TabularTransformationExecutionService()
    spec = _minimal_job_spec("query_result", "inline-1", version="1.1")
    spec["inputs"]["query_result"] = {"rows": [{"country": "CR"}]}
    spec["outputs"]["logical_artifacts"] = [{"name": "report_a", "type": "table"}]
    spec["outputs"]["render"] = [
        {
            "format": "xlsx",
            "options": {
                "sheet_plan": {
                    "sheets": [
                        {"name": "Summary", "source_artifact": "report_a"},
                        {"name": "Summary", "source_artifact": "report_a"},
                    ]
                }
            },
        }
    ]

    with pytest.raises(JobSpecValidationError):
        service.execute(
            context=ReadAuthorizationContext(tenant_id="tenant-1"),
            job_spec=spec,
        )


def test_sheet_plan_semantics_reject_unknown_source_artifact() -> None:
    service = TabularTransformationExecutionService()
    spec = _minimal_job_spec("query_result", "inline-1", version="1.1")
    spec["inputs"]["query_result"] = {"rows": [{"country": "CR"}]}
    spec["outputs"]["logical_artifacts"] = [{"name": "report_a", "type": "table"}]
    spec["outputs"]["render"] = [
        {
            "format": "xlsx",
            "options": {
                "sheet_plan": {
                    "sheets": [{"name": "Summary", "source_artifact": "missing_artifact"}]
                }
            },
        }
    ]

    with pytest.raises(JobSpecValidationError):
        service.execute(
            context=ReadAuthorizationContext(tenant_id="tenant-1"),
            job_spec=spec,
        )


@override_settings(RESOURCE_REGISTRY_TRANSFORMATION_ENGINE_MODE="rules_v1_1")
def test_rules_v11_mode_falls_back_to_legacy_for_v10() -> None:
    service = TabularTransformationExecutionService()
    spec = _minimal_job_spec("query_result", "inline-1", version="1.0")
    spec["inputs"]["query_result"] = {"rows": [{"country": "CR"}]}

    payload = service.execute(
        context=ReadAuthorizationContext(tenant_id="tenant-1"),
        job_spec=spec,
    )

    assert payload["success"] is True
    assert payload["engine_mode"] == "legacy"
