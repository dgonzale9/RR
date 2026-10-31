from __future__ import annotations

from pathlib import Path
from types import SimpleNamespace
from unittest.mock import MagicMock

import duckdb
import pytest

# Registration now reads the persisted organization rollout policy.
pytestmark = pytest.mark.django_db
from django.db import IntegrityError
from openpyxl import Workbook

from resource_registry.application.commands import RegisterDatasetCommand
from resource_registry.application.dataset_registration import DatasetRegistrationFacade
from resource_registry.application.exceptions import (
    DatasetRegistrationError,
    DuplicateDatasetError,
    ValidationError,
)
from resource_registry.application.idempotency import IdempotencyResolution
from resource_registry.application.storage_strategy import DatasetStorageMode, DatasetStoragePlan
from resource_registry.models import Dataset, WorkbookIngestionGroup
from resource_registry.security.policies import ArtifactMaterializationPolicy
from resource_registry.services.registry.dataset_registry_service import (
    VersionContextPersistenceResult,
)
from resource_registry.services.schema_quality_assessor import SchemaQualityAssessment
from resource_registry.storage.base import StorageObjectMetadata

CSV_ROW_COUNT = 3
CSV_COLUMN_COUNT = 3
EXPECTED_DATASET_CREATES = 2
MANAGED_UPLOAD_SIZE_BYTES = 999
EXPECTED_TWO_CALLS = 2


@pytest.fixture(autouse=True)
def allow_local_materialization_policy(monkeypatch: pytest.MonkeyPatch) -> None:
    def _resolve_policy(_self, *, tenant_id: str, dataset_id: str | None = None):
        _ = tenant_id, dataset_id
        return ArtifactMaterializationPolicy(allow_local_files=True)

    def _resolve_with_provenance(_self, *, tenant_id: str, dataset_id: str | None = None):
        _ = tenant_id, dataset_id
        return SimpleNamespace(
            policy=ArtifactMaterializationPolicy(allow_local_files=True),
            provenance=SimpleNamespace(
                policy_key="artifact_materialization_default",
                source="SYSTEM_DEFAULT",
                revision=None,
                rules_hash="hash",
                resolution_path=["tenant_default", "system_default"],
            ),
        )

    monkeypatch.setattr(
        "resource_registry.application.dataset_registration."
        "ArtifactMaterializationPolicyResolver.resolve_policy",
        _resolve_policy,
    )
    monkeypatch.setattr(
        "resource_registry.application.dataset_registration."
        "ArtifactMaterializationPolicyResolver.resolve_with_provenance",
        _resolve_with_provenance,
    )


@pytest.fixture
def csv_dataset_path(tmp_path: Path) -> Path:
    dataset_path = tmp_path / "sample.csv"
    dataset_path.write_text("id,name,score\n1,Ada,95.5\n2,Ben,88.0\n3,Cam,72.5\n", encoding="utf-8")
    return dataset_path


@pytest.fixture
def parquet_dataset_path(tmp_path: Path) -> Path:
    dataset_path = tmp_path / "sample.parquet"
    with duckdb.connect(database=":memory:") as connection:
        connection.execute("""
            CREATE TABLE sample_data AS
            SELECT 1 AS id, 'Ada' AS name, 95.5 AS score
            UNION ALL SELECT 2, 'Ben', 88.0
            UNION ALL SELECT 3, 'Cam', 72.5
            """)
        connection.execute("COPY sample_data TO ? (FORMAT PARQUET)", [str(dataset_path)])
    return dataset_path


@pytest.fixture
def xlsx_dataset_path(tmp_path: Path) -> Path:
    dataset_path = tmp_path / "sample.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Sheet1"
    worksheet.append(["id", "name", "score"])
    worksheet.append([1, "Ada", 95.5])
    worksheet.append([2, "Ben", 88.0])
    worksheet.append([3, "Cam", 72.5])
    workbook.save(dataset_path)
    workbook.close()
    return dataset_path


@pytest.fixture
def xlsx_multi_sheet_path(tmp_path: Path) -> Path:
    dataset_path = tmp_path / "multi.xlsx"
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary.append(["id", "name", "score"])
    summary.append([1, "Ada", 95.5])
    detail = workbook.create_sheet("Detail")
    detail.append(["id", "country", "spend"])
    detail.append([1, "CR", 100.0])
    workbook.save(dataset_path)
    workbook.close()
    return dataset_path


@pytest.fixture
def xlsx_with_empty_sheet_path(tmp_path: Path) -> Path:
    dataset_path = tmp_path / "multi_with_empty.xlsx"
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary.append(["id", "name"])
    summary.append([1, "Ada"])
    workbook.create_sheet("Empty")
    workbook.save(dataset_path)
    workbook.close()
    return dataset_path


@pytest.fixture
def xlsx_all_empty_sheets_path(tmp_path: Path) -> Path:
    dataset_path = tmp_path / "all_empty.xlsx"
    workbook = Workbook()
    first = workbook.active
    first.title = "EmptyOne"
    workbook.create_sheet("EmptyTwo")
    workbook.save(dataset_path)
    workbook.close()
    return dataset_path


def _build_registry_service_mock() -> MagicMock:
    registry_service = MagicMock()
    dataset = SimpleNamespace(
        id="dataset-1",
        versions=SimpleNamespace(
            order_by=lambda *_args, **_kwargs: SimpleNamespace(first=lambda: None),
        ),
    )
    dataset_version = SimpleNamespace(id="version-1")
    dataset_schema = SimpleNamespace(id="schema-1")
    dataset_profile = SimpleNamespace(id="profile-1")
    artifact = SimpleNamespace(id="artifact-1")

    registry_service.create_dataset.return_value = dataset
    registry_service.create_version_bundle_from_context.return_value = (
        VersionContextPersistenceResult(
            dataset_version=dataset_version,
            dataset_schema=dataset_schema,
            dataset_profile=dataset_profile,
        )
    )
    registry_service.attach_artifact.return_value = artifact
    return registry_service


def _build_idempotency_mock(
    *,
    dataset: SimpleNamespace | None = None,
    existing_version: SimpleNamespace | None = None,
    existing_schema: SimpleNamespace | None = None,
    existing_profile: SimpleNamespace | None = None,
    existing_artifact: SimpleNamespace | None = None,
    content_hash: str = "hash-1",
) -> MagicMock:
    idempotency_service = MagicMock()
    idempotency_service.resolve.return_value = IdempotencyResolution(
        dataset=dataset,
        existing_version=existing_version,
        existing_schema=existing_schema,
        existing_profile=existing_profile,
        existing_artifact=existing_artifact,
        content_hash=content_hash,
    )
    if dataset is None:
        idempotency_service.get_next_version_number.return_value = 1
    else:
        idempotency_service.get_next_version_number.return_value = 2
    return idempotency_service


def _build_validator_mock() -> MagicMock:
    validator = MagicMock()
    validator.validate.return_value = None
    return validator


def _build_storage_strategy_mock(
    *,
    should_upload: bool,
    provider: MagicMock | None = None,
    destination_path: str | None = None,
) -> MagicMock:
    strategy = MagicMock()
    mode = (
        DatasetStorageMode.MANAGED_UPLOAD if should_upload else DatasetStorageMode.LOCAL_REFERENCE
    )
    strategy.create_plan.return_value = DatasetStoragePlan(
        mode=mode,
        should_upload=should_upload,
        provider=provider,
        provider_key="azure_blob" if should_upload else None,
        destination_path=destination_path,
    )
    return strategy


def _build_workbook_group_service_mock() -> MagicMock:
    service = MagicMock()
    service.get_or_create_group.return_value = SimpleNamespace(id="workbook-group-1")
    service.mark_completed.return_value = None
    service.mark_failed.return_value = None
    return service


def _build_workbook_job_service_mock() -> MagicMock:
    service = MagicMock()
    service.create_job.return_value = SimpleNamespace(id="job-1", status="queued")
    service.get_job.return_value = SimpleNamespace(
        id="job-1",
        tenant_id="11111111-1111-1111-1111-111111111111",
        workbook_group_id=None,
        status="completed",
        created_at=SimpleNamespace(isoformat=lambda: "2026-01-01T00:00:00"),
        started_at=SimpleNamespace(isoformat=lambda: "2026-01-01T00:00:01"),
        completed_at=SimpleNamespace(isoformat=lambda: "2026-01-01T00:00:02"),
        error_code=None,
        error_message=None,
        retry_count=0,
        result_payload_json={
            "created": [],
            "deduplicated": [],
            "skipped": [],
            "total_processed": 0,
            "total_created": 0,
            "total_deduplicated": 0,
            "total_skipped": 0,
        },
    )
    return service


def test_register_dataset_from_csv_path(csv_dataset_path: Path) -> None:
    registry_service = _build_registry_service_mock()
    facade = DatasetRegistrationFacade(
        registry_service=registry_service,
        validator=_build_validator_mock(),
        idempotency_service=_build_idempotency_mock(),
    )
    command = RegisterDatasetCommand(
        tenant_id="11111111-1111-1111-1111-111111111111",
        slug="sales-data",
        name="Sales Data",
        path=str(csv_dataset_path),
        source_type="FILE_UPLOAD",
    )

    result = facade.register(command)

    assert result.dataset_id == "dataset-1"
    assert result.dataset_version_id == "version-1"
    assert result.schema_id == "schema-1"
    assert result.profile_id == "profile-1"
    assert result.artifact_id == "artifact-1"
    assert result.row_count == CSV_ROW_COUNT
    assert result.column_count == CSV_COLUMN_COUNT
    assert result.source_format == "csv"
    assert result.file_size_bytes > 0
    assert registry_service.attach_artifact.call_count == EXPECTED_TWO_CALLS
    canonical_call = registry_service.attach_artifact.call_args_list[1].kwargs
    assert canonical_call["format"] == "PARQUET"
    assert canonical_call["metadata_json"]["conversion_applied"] is True


def test_register_dataset_from_parquet_path(parquet_dataset_path: Path) -> None:
    registry_service = _build_registry_service_mock()
    facade = DatasetRegistrationFacade(
        registry_service=registry_service,
        validator=_build_validator_mock(),
        idempotency_service=_build_idempotency_mock(),
    )
    command = RegisterDatasetCommand(
        tenant_id="11111111-1111-1111-1111-111111111111",
        slug="sales-data-pq",
        name="Sales Data PQ",
        path=str(parquet_dataset_path),
        source_type="FILE_UPLOAD",
    )

    result = facade.register(command)

    assert result.source_format == "parquet"
    assert result.row_count == CSV_ROW_COUNT
    assert result.column_count == CSV_COLUMN_COUNT
    assert result.file_size_bytes > 0
    assert registry_service.attach_artifact.call_count == EXPECTED_TWO_CALLS
    canonical_call = registry_service.attach_artifact.call_args_list[1].kwargs
    assert canonical_call["format"] == "PARQUET"
    assert canonical_call["metadata_json"]["conversion_applied"] is False


def test_register_dataset_from_xlsx_path(xlsx_dataset_path: Path) -> None:
    registry_service = _build_registry_service_mock()
    facade = DatasetRegistrationFacade(
        registry_service=registry_service,
        validator=_build_validator_mock(),
        idempotency_service=_build_idempotency_mock(),
    )
    command = RegisterDatasetCommand(
        tenant_id="11111111-1111-1111-1111-111111111111",
        slug="sales-data-xlsx",
        name="Sales Data XLSX",
        path=str(xlsx_dataset_path),
        source_type="FILE_UPLOAD",
    )

    result = facade.register(command)

    assert result.source_format == "xlsx"
    assert result.row_count == CSV_ROW_COUNT
    assert result.column_count == CSV_COLUMN_COUNT
    assert result.file_size_bytes > 0
    assert registry_service.attach_artifact.call_count == EXPECTED_TWO_CALLS
    canonical_call = registry_service.attach_artifact.call_args_list[1].kwargs
    assert canonical_call["format"] == "PARQUET"
    assert canonical_call["metadata_json"]["excel_sheet_name"] == "Sheet1"


def test_register_dataset_from_xlsx_with_sheet_name(xlsx_dataset_path: Path) -> None:
    registry_service = _build_registry_service_mock()
    facade = DatasetRegistrationFacade(
        registry_service=registry_service,
        validator=_build_validator_mock(),
        idempotency_service=_build_idempotency_mock(),
    )
    command = RegisterDatasetCommand(
        tenant_id="11111111-1111-1111-1111-111111111111",
        slug="sales-data-xlsx-sheet",
        name="Sales Data XLSX Sheet",
        path=str(xlsx_dataset_path),
        source_type="FILE_UPLOAD",
        sheet_name="Sheet1",
    )

    result = facade.register(command)

    assert result.source_format == "xlsx"
    canonical_call = registry_service.attach_artifact.call_args_list[1].kwargs
    assert canonical_call["metadata_json"]["excel_sheet_name"] == "Sheet1"


def test_register_dataset_from_xlsx_with_missing_sheet_returns_stable_error(
    xlsx_dataset_path: Path,
) -> None:
    facade = DatasetRegistrationFacade(
        registry_service=_build_registry_service_mock(),
        validator=_build_validator_mock(),
        idempotency_service=_build_idempotency_mock(),
    )
    command = RegisterDatasetCommand(
        tenant_id="11111111-1111-1111-1111-111111111111",
        slug="sales-data-xlsx-missing-sheet",
        name="Sales Data XLSX Missing Sheet",
        path=str(xlsx_dataset_path),
        source_type="FILE_UPLOAD",
        sheet_name="MissingSheet",
    )

    with pytest.raises(ValidationError) as error:
        facade.register(command)
    assert error.value.code == "worksheet_not_found"
    assert error.value.status == 400
    assert str(error.value) == "The requested Excel worksheet was not found."


def test_register_workbook_ingests_all_sheets_and_derives_slug_suffixes(
    xlsx_multi_sheet_path: Path,
) -> None:
    registry_service = _build_registry_service_mock()
    facade = DatasetRegistrationFacade(
        registry_service=registry_service,
        validator=_build_validator_mock(),
        idempotency_service=_build_idempotency_mock(),
        workbook_group_service=_build_workbook_group_service_mock(),
    )
    command = RegisterDatasetCommand(
        tenant_id="11111111-1111-1111-1111-111111111111",
        slug="sales-data-workbook",
        name="Sales Data Workbook",
        path=str(xlsx_multi_sheet_path),
        source_type="FILE_UPLOAD",
    )

    result = facade.register_workbook(command)

    assert result.total_processed == 2
    assert result.total_created == 2
    assert result.total_deduplicated == 0
    assert result.total_skipped == 0
    created_slugs = [call.kwargs["slug"] for call in registry_service.create_dataset.call_args_list]
    assert created_slugs == ["sales-data-workbook-summary", "sales-data-workbook-detail"]
    created_names = [call.kwargs["name"] for call in registry_service.create_dataset.call_args_list]
    assert created_names == ["Sales Data Workbook — Summary", "Sales Data Workbook — Detail"]
    assert len(set(created_names)) == len(created_names)
    attached_sheet_names = [
        call.kwargs["sheet_name"] for call in registry_service.attach_artifact.call_args_list
    ]
    assert attached_sheet_names == ["Summary", "Summary", "Detail", "Detail"]
    assert registry_service.attach_artifact.call_count == EXPECTED_TWO_CALLS * EXPECTED_TWO_CALLS


def test_register_workbook_uses_workbook_name_for_single_sheet_xlsx(
    xlsx_dataset_path: Path,
) -> None:
    registry_service = _build_registry_service_mock()
    facade = DatasetRegistrationFacade(
        registry_service=registry_service,
        validator=_build_validator_mock(),
        idempotency_service=_build_idempotency_mock(),
        workbook_group_service=_build_workbook_group_service_mock(),
    )
    command = RegisterDatasetCommand(
        tenant_id="11111111-1111-1111-1111-111111111111",
        slug="customer-data",
        name="Customer Data",
        path=str(xlsx_dataset_path),
        source_type="FILE_UPLOAD",
    )

    result = facade.register_workbook(command)

    assert result.total_processed == 1
    assert result.total_created == 1
    assert registry_service.create_dataset.call_args.kwargs["name"] == "Customer Data"
    assert registry_service.attach_artifact.call_args_list[0].kwargs["sheet_name"] == "Sheet1"


def test_register_workbook_skips_invalid_sheets_and_continues(
    xlsx_with_empty_sheet_path: Path,
) -> None:
    registry_service = _build_registry_service_mock()
    facade = DatasetRegistrationFacade(
        registry_service=registry_service,
        validator=_build_validator_mock(),
        idempotency_service=_build_idempotency_mock(),
        workbook_group_service=_build_workbook_group_service_mock(),
    )
    command = RegisterDatasetCommand(
        tenant_id="11111111-1111-1111-1111-111111111111",
        slug="sales-data-workbook",
        name="Sales Data Workbook",
        path=str(xlsx_with_empty_sheet_path),
        source_type="FILE_UPLOAD",
    )

    result = facade.register_workbook(command)

    assert result.total_processed == 2
    assert result.total_created == 1
    assert result.total_deduplicated == 0
    assert result.total_skipped == 1
    assert result.skipped[0].sheet_name == "Empty"
    assert result.skipped[0].code == "empty_dataset"


def test_register_workbook_fails_when_all_sheets_are_invalid(
    xlsx_all_empty_sheets_path: Path,
) -> None:
    facade = DatasetRegistrationFacade(
        registry_service=_build_registry_service_mock(),
        validator=_build_validator_mock(),
        idempotency_service=_build_idempotency_mock(),
        workbook_group_service=_build_workbook_group_service_mock(),
    )
    command = RegisterDatasetCommand(
        tenant_id="11111111-1111-1111-1111-111111111111",
        slug="sales-data-workbook",
        name="Sales Data Workbook",
        path=str(xlsx_all_empty_sheets_path),
        source_type="FILE_UPLOAD",
    )

    with pytest.raises(ValidationError) as error:
        facade.register_workbook(command)
    assert error.value.code == "workbook_has_no_valid_sheets"
    assert error.value.status == 400


def test_register_workbook_reports_deduplicated_and_created_results(
    xlsx_multi_sheet_path: Path,
) -> None:
    existing_dataset = SimpleNamespace(id="dataset-existing")
    existing_version = SimpleNamespace(id="version-existing")
    existing_schema = SimpleNamespace(id="schema-existing")
    existing_profile = SimpleNamespace(id="profile-existing")
    existing_artifact = SimpleNamespace(id="artifact-existing")

    def _resolve(
        tenant_id: str,
        slug: str,
        path: str,
        precomputed_content_hash: str | None = None,
    ):
        _ = tenant_id, path
        _ = precomputed_content_hash
        if slug.endswith("-summary"):
            return IdempotencyResolution(
                dataset=existing_dataset,
                existing_version=existing_version,
                existing_schema=existing_schema,
                existing_profile=existing_profile,
                existing_artifact=existing_artifact,
                content_hash="hash-existing",
            )
        return IdempotencyResolution(
            dataset=None,
            existing_version=None,
            existing_schema=None,
            existing_profile=None,
            existing_artifact=None,
            content_hash="hash-new",
        )

    idempotency_service = MagicMock()
    idempotency_service.resolve.side_effect = _resolve
    idempotency_service.get_next_version_number.return_value = 1

    registry_service = _build_registry_service_mock()
    facade = DatasetRegistrationFacade(
        registry_service=registry_service,
        validator=_build_validator_mock(),
        idempotency_service=idempotency_service,
        workbook_group_service=_build_workbook_group_service_mock(),
    )
    command = RegisterDatasetCommand(
        tenant_id="11111111-1111-1111-1111-111111111111",
        slug="sales-data-workbook",
        name="Sales Data Workbook",
        path=str(xlsx_multi_sheet_path),
        source_type="FILE_UPLOAD",
    )

    result = facade.register_workbook(command)

    assert result.total_created == 1
    assert result.total_deduplicated == 1
    assert result.total_skipped == 0


def test_register_workbook_keeps_single_result_behavior_for_non_xlsx(
    csv_dataset_path: Path,
) -> None:
    registry_service = _build_registry_service_mock()
    facade = DatasetRegistrationFacade(
        registry_service=registry_service,
        validator=_build_validator_mock(),
        idempotency_service=_build_idempotency_mock(),
        workbook_group_service=_build_workbook_group_service_mock(),
    )
    command = RegisterDatasetCommand(
        tenant_id="11111111-1111-1111-1111-111111111111",
        slug="sales-data",
        name="Sales Data",
        path=str(csv_dataset_path),
        source_type="FILE_UPLOAD",
    )

    result = facade.register_workbook(command)

    assert result.total_processed == 1
    assert result.total_created == 1
    assert len(result.created) == 1
    assert result.created[0].source_format == "csv"


@pytest.mark.django_db
def test_register_workbook_creates_workbook_group_and_links_datasets(
    xlsx_multi_sheet_path: Path,
) -> None:
    facade = DatasetRegistrationFacade()
    command = RegisterDatasetCommand(
        tenant_id="11111111-1111-1111-1111-111111111111",
        slug="workbook-group-link",
        name="Workbook Group Link",
        path=str(xlsx_multi_sheet_path),
        source_type="FILE_UPLOAD",
        storage_mode="LOCAL_REFERENCE",
    )

    result = facade.register_workbook(command)

    assert result.total_created == 2
    group = WorkbookIngestionGroup.objects.get(
        tenant_id=command.tenant_id,
        source_hash=facade._idempotency_service.compute_content_hash(command.path),
    )
    assert group.total_sheets == 2
    assert group.total_created == 2
    assert group.total_skipped == 0
    assert Dataset.objects.filter(workbook_group=group).count() == 2


def test_register_workbook_reuses_precomputed_hash_for_idempotency(
    xlsx_multi_sheet_path: Path,
) -> None:
    idempotency_service = MagicMock()
    idempotency_service.compute_content_hash.return_value = "hash-123"
    idempotency_service.resolve.return_value = IdempotencyResolution(
        dataset=None,
        existing_version=None,
        existing_schema=None,
        existing_profile=None,
        existing_artifact=None,
        content_hash="hash-123",
    )
    idempotency_service.get_next_version_number.return_value = 1
    facade = DatasetRegistrationFacade(
        registry_service=_build_registry_service_mock(),
        validator=_build_validator_mock(),
        idempotency_service=idempotency_service,
        workbook_group_service=_build_workbook_group_service_mock(),
    )
    command = RegisterDatasetCommand(
        tenant_id="11111111-1111-1111-1111-111111111111",
        slug="reuse-hash",
        name="Reuse Hash",
        path=str(xlsx_multi_sheet_path),
        source_type="FILE_UPLOAD",
    )

    facade.register_workbook(command)

    assert idempotency_service.compute_content_hash.call_count == 1
    for call in idempotency_service.resolve.call_args_list:
        assert call.kwargs["precomputed_content_hash"] == "hash-123"


def test_start_workbook_ingestion_creates_queued_job(
    xlsx_multi_sheet_path: Path,
) -> None:
    workbook_job_service = _build_workbook_job_service_mock()
    facade = DatasetRegistrationFacade(
        registry_service=_build_registry_service_mock(),
        validator=_build_validator_mock(),
        idempotency_service=_build_idempotency_mock(),
        workbook_job_service=workbook_job_service,
    )
    command = RegisterDatasetCommand(
        tenant_id="11111111-1111-1111-1111-111111111111",
        slug="async-workbook",
        name="Async Workbook",
        path=str(xlsx_multi_sheet_path),
        source_type="FILE_UPLOAD",
    )

    result = facade.start_workbook_ingestion(command)

    assert result.ingestion_job_id == "job-1"
    assert result.status == "queued"
    workbook_job_service.create_job.assert_called_once()


def test_get_workbook_ingestion_status_returns_status_payload() -> None:
    workbook_job_service = _build_workbook_job_service_mock()
    facade = DatasetRegistrationFacade(
        workbook_job_service=workbook_job_service,
        validator=_build_validator_mock(),
    )

    result = facade.get_workbook_ingestion_status(
        tenant_id="11111111-1111-1111-1111-111111111111",
        ingestion_job_id="job-1",
    )

    assert result.ingestion_job_id == "job-1"
    assert result.status == "completed"
    assert result.retry_count == 0


def test_get_workbook_ingestion_result_returns_batch_payload() -> None:
    workbook_job_service = _build_workbook_job_service_mock()
    facade = DatasetRegistrationFacade(
        workbook_job_service=workbook_job_service,
        validator=_build_validator_mock(),
    )

    result = facade.get_workbook_ingestion_result(
        tenant_id="11111111-1111-1111-1111-111111111111",
        ingestion_job_id="job-1",
    )

    assert result.ingestion_job_id == "job-1"
    assert result.status == "completed"
    assert result.result is not None
    assert result.result.total_processed == 0


def test_preview_rows_are_not_persisted_anywhere(csv_dataset_path: Path) -> None:
    registry_service = _build_registry_service_mock()
    facade = DatasetRegistrationFacade(
        registry_service=registry_service,
        validator=_build_validator_mock(),
        idempotency_service=_build_idempotency_mock(),
    )
    command = RegisterDatasetCommand(
        tenant_id="11111111-1111-1111-1111-111111111111",
        slug="sales-data-preview",
        name="Sales Data Preview",
        path=str(csv_dataset_path),
        source_type="FILE_UPLOAD",
    )

    facade.register(command)

    artifact_metadata = registry_service.attach_artifact.call_args.kwargs["metadata_json"]
    assert "preview" not in artifact_metadata
    persisted_context = registry_service.create_version_bundle_from_context.call_args.kwargs[
        "context"
    ]
    assert "preview" in persisted_context


def test_canonical_artifact_is_persisted_with_normalization_metadata(
    csv_dataset_path: Path,
) -> None:
    registry_service = _build_registry_service_mock()
    facade = DatasetRegistrationFacade(
        registry_service=registry_service,
        validator=_build_validator_mock(),
        idempotency_service=_build_idempotency_mock(),
    )
    command = RegisterDatasetCommand(
        tenant_id="11111111-1111-1111-1111-111111111111",
        slug="canonical-sales",
        name="Canonical Sales",
        path=str(csv_dataset_path),
        source_type="FILE_UPLOAD",
    )

    facade.register(command)

    assert registry_service.attach_artifact.call_count == EXPECTED_TWO_CALLS
    source_call = registry_service.attach_artifact.call_args_list[0].kwargs
    canonical_call = registry_service.attach_artifact.call_args_list[1].kwargs
    assert source_call["artifact_role"] == "PRIMARY"
    assert canonical_call["artifact_role"] == "SNAPSHOT"
    assert canonical_call["format"] == "PARQUET"
    assert canonical_call["metadata_json"]["is_canonical"] is True
    assert canonical_call["metadata_json"]["source_artifact_id"] == "artifact-1"
    assert canonical_call["metadata_json"]["normalization_version"] == "v1"


def test_invalid_path_fails_before_partial_writes() -> None:
    registry_service = _build_registry_service_mock()
    facade = DatasetRegistrationFacade(registry_service=registry_service)
    command = RegisterDatasetCommand(
        tenant_id="11111111-1111-1111-1111-111111111111",
        slug="invalid-path",
        name="Invalid Path",
        path="/tmp/does-not-exist.csv",
        source_type="FILE_UPLOAD",
    )

    with pytest.raises(ValidationError):
        facade.register(command)
    registry_service.create_dataset.assert_not_called()


def test_duplicate_dataset_slug_for_same_tenant_is_rejected(csv_dataset_path: Path) -> None:
    registry_service = _build_registry_service_mock()
    registry_service.create_dataset.side_effect = IntegrityError("duplicate key value")
    facade = DatasetRegistrationFacade(
        registry_service=registry_service,
        validator=_build_validator_mock(),
        idempotency_service=_build_idempotency_mock(),
    )
    command = RegisterDatasetCommand(
        tenant_id="11111111-1111-1111-1111-111111111111",
        slug="same-slug",
        name="Dataset 1",
        path=str(csv_dataset_path),
        source_type="FILE_UPLOAD",
    )

    with pytest.raises(DuplicateDatasetError):
        facade.register(command)


def test_same_slug_for_different_tenant_is_allowed(csv_dataset_path: Path) -> None:
    registry_service = _build_registry_service_mock()
    facade = DatasetRegistrationFacade(
        registry_service=registry_service,
        validator=_build_validator_mock(),
        idempotency_service=_build_idempotency_mock(),
    )
    command_a = RegisterDatasetCommand(
        tenant_id="11111111-1111-1111-1111-111111111111",
        slug="shared-slug",
        name="Tenant A Dataset",
        path=str(csv_dataset_path),
        source_type="FILE_UPLOAD",
    )
    command_b = RegisterDatasetCommand(
        tenant_id="22222222-2222-2222-2222-222222222222",
        slug="shared-slug",
        name="Tenant B Dataset",
        path=str(csv_dataset_path),
        source_type="FILE_UPLOAD",
    )

    facade.register(command_a)
    facade.register(command_b)

    assert registry_service.create_dataset.call_count == EXPECTED_DATASET_CREATES
    first_tenant = registry_service.create_dataset.call_args_list[0].kwargs["tenant_id"]
    second_tenant = registry_service.create_dataset.call_args_list[1].kwargs["tenant_id"]
    assert first_tenant != second_tenant


def test_facade_composes_context_builder_and_registry_service(csv_dataset_path: Path) -> None:
    context_builder = MagicMock()
    context_builder.build_context.return_value = {
        "source": {
            "path": str(csv_dataset_path),
            "format": "csv",
            "file_name": "sample.csv",
            "file_size_bytes": 100,
        },
        "summary": {
            "row_count": CSV_ROW_COUNT,
            "column_count": CSV_COLUMN_COUNT,
        },
        "schema": [{"name": "id", "type": "BIGINT"}],
        "preview": [{"id": 1}],
        "safety": {"preview_limit": 10, "preview_limited": True, "raw_content_included": False},
    }
    registry_service = _build_registry_service_mock()
    validator = _build_validator_mock()
    idempotency_service = _build_idempotency_mock()
    facade = DatasetRegistrationFacade(
        context_builder=context_builder,
        registry_service=registry_service,
        validator=validator,
        idempotency_service=idempotency_service,
    )
    command = RegisterDatasetCommand(
        tenant_id="11111111-1111-1111-1111-111111111111",
        slug="compose-check",
        name="Compose Check",
        path=str(csv_dataset_path),
        source_type="FILE_UPLOAD",
    )

    facade.register(command)

    context_builder.build_context.assert_called_once_with(
        str(csv_dataset_path),
        preview_limit=10,
        sheet_name=None,
        mime_type=None,
    )
    validator.validate.assert_called_once()
    idempotency_service.resolve.assert_called_once()
    registry_service.create_dataset.assert_called_once()
    registry_service.create_version_bundle_from_context.assert_called_once()


def test_low_confidence_schema_fails_by_default(csv_dataset_path: Path) -> None:
    schema_quality_assessor = MagicMock()
    schema_quality_assessor.assess.return_value = SchemaQualityAssessment(
        score=0.20,
        confidence="low",
        warnings=("Detected many generic column names (column_N).",),
        metrics={},
    )
    facade = DatasetRegistrationFacade(
        registry_service=_build_registry_service_mock(),
        validator=_build_validator_mock(),
        idempotency_service=_build_idempotency_mock(),
        schema_quality_assessor=schema_quality_assessor,
    )
    command = RegisterDatasetCommand(
        tenant_id="11111111-1111-1111-1111-111111111111",
        slug="low-confidence-fail",
        name="Low Confidence Fail",
        path=str(csv_dataset_path),
        source_type="FILE_UPLOAD",
    )

    with pytest.raises(ValidationError) as error:
        facade.register(command)
    assert error.value.code == "low_confidence_schema"


def test_low_confidence_schema_can_be_overridden(csv_dataset_path: Path) -> None:
    schema_quality_assessor = MagicMock()
    schema_quality_assessor.assess.return_value = SchemaQualityAssessment(
        score=0.20,
        confidence="low",
        warnings=("Detected many generic column names (column_N).",),
        metrics={"generic_ratio": 0.5},
    )
    registry_service = _build_registry_service_mock()
    facade = DatasetRegistrationFacade(
        registry_service=registry_service,
        validator=_build_validator_mock(),
        idempotency_service=_build_idempotency_mock(),
        schema_quality_assessor=schema_quality_assessor,
    )
    command = RegisterDatasetCommand(
        tenant_id="11111111-1111-1111-1111-111111111111",
        slug="low-confidence-override",
        name="Low Confidence Override",
        path=str(csv_dataset_path),
        source_type="FILE_UPLOAD",
        allow_low_confidence_schema=True,
    )

    result = facade.register(command)

    assert result.dataset_id == "dataset-1"
    source_artifact_call = registry_service.attach_artifact.call_args_list[0].kwargs
    assert source_artifact_call["metadata_json"]["schema_quality"]["confidence"] == "low"


def test_same_tenant_slug_same_hash_returns_existing_result(csv_dataset_path: Path) -> None:
    existing_dataset = SimpleNamespace(id="dataset-1")
    existing_version = SimpleNamespace(id="version-existing")
    existing_schema = SimpleNamespace(id="schema-existing")
    existing_profile = SimpleNamespace(id="profile-existing")
    existing_artifact = SimpleNamespace(id="artifact-existing")
    registry_service = _build_registry_service_mock()
    idempotency_service = _build_idempotency_mock(
        dataset=existing_dataset,
        existing_version=existing_version,
        existing_schema=existing_schema,
        existing_profile=existing_profile,
        existing_artifact=existing_artifact,
    )
    facade = DatasetRegistrationFacade(
        registry_service=registry_service,
        validator=_build_validator_mock(),
        idempotency_service=idempotency_service,
    )
    command = RegisterDatasetCommand(
        tenant_id="11111111-1111-1111-1111-111111111111",
        slug="same-hash",
        name="Same Hash",
        path=str(csv_dataset_path),
        source_type="FILE_UPLOAD",
    )

    result = facade.register(command)

    assert result.dataset_version_id == "version-existing"
    registry_service.create_version_bundle_from_context.assert_not_called()
    registry_service.attach_artifact.assert_not_called()


def test_same_tenant_slug_different_hash_creates_new_version(csv_dataset_path: Path) -> None:
    existing_dataset = SimpleNamespace(
        id="dataset-1",
        versions=SimpleNamespace(
            order_by=lambda *_args, **_kwargs: SimpleNamespace(first=lambda: None),
        ),
    )
    registry_service = _build_registry_service_mock()
    registry_service.create_dataset.reset_mock()
    idempotency_service = _build_idempotency_mock(dataset=existing_dataset, content_hash="new-hash")
    facade = DatasetRegistrationFacade(
        registry_service=registry_service,
        validator=_build_validator_mock(),
        idempotency_service=idempotency_service,
    )
    command = RegisterDatasetCommand(
        tenant_id="11111111-1111-1111-1111-111111111111",
        slug="different-hash",
        name="Different Hash",
        path=str(csv_dataset_path),
        source_type="FILE_UPLOAD",
    )

    facade.register(command)

    registry_service.create_dataset.assert_not_called()
    registry_service.create_version_bundle_from_context.assert_called_once()


def test_application_validation_error_is_propagated(csv_dataset_path: Path) -> None:
    validator = MagicMock()
    validator.validate.side_effect = ValidationError("invalid")
    facade = DatasetRegistrationFacade(
        registry_service=_build_registry_service_mock(),
        validator=validator,
        idempotency_service=_build_idempotency_mock(),
    )
    command = RegisterDatasetCommand(
        tenant_id="11111111-1111-1111-1111-111111111111",
        slug="validation",
        name="Validation",
        path=str(csv_dataset_path),
        source_type="FILE_UPLOAD",
    )

    with pytest.raises(ValidationError):
        facade.register(command)


def test_upload_invoked_when_managed_upload_enabled(csv_dataset_path: Path) -> None:
    registry_service = _build_registry_service_mock()
    provider = MagicMock()
    artifact_lifecycle_service = MagicMock()
    artifact_lifecycle_service.register_uploaded_artifact.return_value = SimpleNamespace(
        id="artifact-1",
    )
    provider.put.return_value = StorageObjectMetadata(
        uri="azure://media/datasets/tenant/sales-data/v1/original.csv",
        provider="azure_blob",
        size_bytes=MANAGED_UPLOAD_SIZE_BYTES,
        checksum="hash-1",
        metadata={"safe": "ok"},
    )
    facade = DatasetRegistrationFacade(
        registry_service=registry_service,
        validator=_build_validator_mock(),
        idempotency_service=_build_idempotency_mock(),
        storage_strategy=_build_storage_strategy_mock(
            should_upload=True,
            provider=provider,
            destination_path="datasets/tenant/sales-data/v1/original.csv",
        ),
        artifact_lifecycle_service=artifact_lifecycle_service,
    )
    command = RegisterDatasetCommand(
        tenant_id="11111111-1111-1111-1111-111111111111",
        slug="sales-data",
        name="Sales Data",
        path=str(csv_dataset_path),
        source_type="FILE_UPLOAD",
        storage_mode="MANAGED_UPLOAD",
        storage_provider="AZURE_BLOB",
    )

    facade.register(command)

    assert provider.put.call_count == EXPECTED_TWO_CALLS
    assert provider.put.call_args_list[0].args == (
        str(csv_dataset_path),
        "datasets/tenant/sales-data/v1/original.csv",
    )
    artifact_kwargs = artifact_lifecycle_service.register_uploaded_artifact.call_args.kwargs
    assert artifact_kwargs["uri"] == "azure://media/datasets/tenant/sales-data/v1/original.csv"
    assert artifact_kwargs["storage_provider"] == "azure_blob"
    assert artifact_kwargs["size_bytes"] == MANAGED_UPLOAD_SIZE_BYTES
    assert artifact_kwargs["checksum"] == "hash-1"


def test_upload_skipped_when_local_reference_mode(csv_dataset_path: Path) -> None:
    registry_service = _build_registry_service_mock()
    strategy = _build_storage_strategy_mock(should_upload=False)
    facade = DatasetRegistrationFacade(
        registry_service=registry_service,
        validator=_build_validator_mock(),
        idempotency_service=_build_idempotency_mock(),
        storage_strategy=strategy,
    )
    command = RegisterDatasetCommand(
        tenant_id="11111111-1111-1111-1111-111111111111",
        slug="local-ref",
        name="Local Ref",
        path=str(csv_dataset_path),
        source_type="FILE_UPLOAD",
        storage_mode="LOCAL_REFERENCE",
    )

    facade.register(command)

    strategy.create_plan.assert_called_once()


def test_duplicate_registration_skips_upload(csv_dataset_path: Path) -> None:
    existing_dataset = SimpleNamespace(id="dataset-1")
    existing_version = SimpleNamespace(id="version-existing")
    existing_schema = SimpleNamespace(id="schema-existing")
    existing_profile = SimpleNamespace(id="profile-existing")
    existing_artifact = SimpleNamespace(id="artifact-existing")
    strategy = MagicMock()
    facade = DatasetRegistrationFacade(
        registry_service=_build_registry_service_mock(),
        validator=_build_validator_mock(),
        idempotency_service=_build_idempotency_mock(
            dataset=existing_dataset,
            existing_version=existing_version,
            existing_schema=existing_schema,
            existing_profile=existing_profile,
            existing_artifact=existing_artifact,
        ),
        storage_strategy=strategy,
    )
    command = RegisterDatasetCommand(
        tenant_id="11111111-1111-1111-1111-111111111111",
        slug="same-hash",
        name="Same Hash",
        path=str(csv_dataset_path),
        source_type="FILE_UPLOAD",
        storage_mode="MANAGED_UPLOAD",
    )

    facade.register(command)

    strategy.create_plan.assert_not_called()


def test_same_slug_different_hash_uploads_new_version(csv_dataset_path: Path) -> None:
    existing_dataset = SimpleNamespace(
        id="dataset-1",
        versions=SimpleNamespace(
            order_by=lambda *_args, **_kwargs: SimpleNamespace(first=lambda: None),
        ),
    )
    registry_service = _build_registry_service_mock()
    provider = MagicMock()
    artifact_lifecycle_service = MagicMock()
    artifact_lifecycle_service.register_uploaded_artifact.return_value = SimpleNamespace(
        id="artifact-1",
    )
    provider.put.return_value = StorageObjectMetadata(
        uri="azure://media/datasets/tenant/different-hash/v2/original.csv",
        provider="azure_blob",
        size_bytes=111,
        checksum="hash-2",
    )
    facade = DatasetRegistrationFacade(
        registry_service=registry_service,
        validator=_build_validator_mock(),
        idempotency_service=_build_idempotency_mock(
            dataset=existing_dataset,
            content_hash="new-hash",
        ),
        storage_strategy=_build_storage_strategy_mock(
            should_upload=True,
            provider=provider,
            destination_path="datasets/tenant/different-hash/v2/original.csv",
        ),
        artifact_lifecycle_service=artifact_lifecycle_service,
    )
    command = RegisterDatasetCommand(
        tenant_id="11111111-1111-1111-1111-111111111111",
        slug="different-hash",
        name="Different Hash",
        path=str(csv_dataset_path),
        source_type="FILE_UPLOAD",
        storage_mode="MANAGED_UPLOAD",
        storage_provider="AZURE_BLOB",
    )

    facade.register(command)

    assert provider.put.call_count == EXPECTED_TWO_CALLS
    registry_service.create_version_bundle_from_context.assert_called_once()


def test_unsafe_storage_uri_is_rejected_before_writes(csv_dataset_path: Path) -> None:
    registry_service = _build_registry_service_mock()
    provider = MagicMock()
    provider.put.return_value = StorageObjectMetadata(
        uri="https://example.blob.core.windows.net/c/file.csv?sig=secret",
        provider="azure_blob",
        size_bytes=100,
    )
    facade = DatasetRegistrationFacade(
        registry_service=registry_service,
        validator=_build_validator_mock(),
        idempotency_service=_build_idempotency_mock(),
        storage_strategy=_build_storage_strategy_mock(
            should_upload=True,
            provider=provider,
            destination_path="datasets/tenant/sales-data/v1/original.csv",
        ),
    )
    command = RegisterDatasetCommand(
        tenant_id="11111111-1111-1111-1111-111111111111",
        slug="sales-data",
        name="Sales Data",
        path=str(csv_dataset_path),
        source_type="FILE_UPLOAD",
        storage_mode="MANAGED_UPLOAD",
    )

    with pytest.raises(DatasetRegistrationError):
        facade.register(command)

    registry_service.create_dataset.assert_not_called()


def test_managed_upload_persistence_failure_triggers_orphan_cleanup(
    csv_dataset_path: Path,
) -> None:
    registry_service = _build_registry_service_mock()
    provider = MagicMock()
    provider.put.return_value = StorageObjectMetadata(
        uri="azure://media/datasets/tenant/sales-data/v1/original.csv",
        provider="azure_blob",
        size_bytes=MANAGED_UPLOAD_SIZE_BYTES,
        checksum="hash-1",
    )
    artifact_lifecycle_service = MagicMock()
    artifact_lifecycle_service.register_uploaded_artifact.side_effect = IntegrityError(
        "artifact write failed",
    )
    facade = DatasetRegistrationFacade(
        registry_service=registry_service,
        validator=_build_validator_mock(),
        idempotency_service=_build_idempotency_mock(),
        storage_strategy=_build_storage_strategy_mock(
            should_upload=True,
            provider=provider,
            destination_path="datasets/tenant/sales-data/v1/original.csv",
        ),
        artifact_lifecycle_service=artifact_lifecycle_service,
    )
    command = RegisterDatasetCommand(
        tenant_id="11111111-1111-1111-1111-111111111111",
        slug="sales-data",
        name="Sales Data",
        path=str(csv_dataset_path),
        source_type="FILE_UPLOAD",
        storage_mode="MANAGED_UPLOAD",
        storage_provider="AZURE_BLOB",
    )

    with pytest.raises(DuplicateDatasetError) as excinfo:
        facade.register(command)

    assert isinstance(excinfo.value.__cause__, IntegrityError)
    artifact_lifecycle_service.cleanup_orphan_upload.assert_called_once()
    cleanup_kwargs = artifact_lifecycle_service.cleanup_orphan_upload.call_args.kwargs
    assert cleanup_kwargs["tenant_id"] == command.tenant_id
    assert cleanup_kwargs["uri"] == "azure://media/datasets/tenant/sales-data/v1/original.csv"


def test_local_registration_rejected_when_policy_disallows_local_files(
    csv_dataset_path: Path,
) -> None:
    resolver = MagicMock()
    resolver.resolve_policy.return_value = ArtifactMaterializationPolicy(
        allow_local_files=False,
    )
    facade = DatasetRegistrationFacade(
        registry_service=_build_registry_service_mock(),
        validator=_build_validator_mock(),
        idempotency_service=_build_idempotency_mock(),
        materialization_policy_resolver=resolver,
    )
    command = RegisterDatasetCommand(
        tenant_id="11111111-1111-1111-1111-111111111111",
        slug="blocked-local",
        name="Blocked Local",
        path=str(csv_dataset_path),
        source_type="FILE_UPLOAD",
    )

    with pytest.raises(DatasetRegistrationError):
        facade.register(command)
