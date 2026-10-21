from __future__ import annotations

import base64
import io
import logging
import re
from dataclasses import dataclass, field
from typing import Any, Protocol

from django.conf import settings
from jsonschema import Draft7Validator
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from resource_registry.read_auth.context import ReadAuthorizationContext
from resource_registry.read_auth.policy import AuthorizedDataRegistryReadFacade
from resource_registry.services.dataset_query_executor import DatasetQueryExecutor
from resource_registry.services.payload_sanitizer import sanitize_payload

_logger = logging.getLogger("app-logger")

_DEFAULT_TRANSFORMATION_ENGINE_MODE = "legacy"
_SUPPORTED_ENGINE_MODES = {"legacy", "rules_v1_1"}


_JOB_SPEC_SCHEMA_V1: dict[str, Any] = {
    "$schema": "http://json-schema.org/draft-07/schema#",
    "type": "object",
    "additionalProperties": False,
    "required": ["job_type", "version", "inputs", "pipeline", "outputs", "constraints"],
    "properties": {
        "job_type": {"type": "string", "const": "tabular_transformation"},
        "version": {"type": "string", "enum": ["1.0", "1.1"]},
        "inputs": {
            "type": "object",
            "additionalProperties": False,
            "required": ["scope"],
            "properties": {
                "scope": {
                    "type": "object",
                    "additionalProperties": False,
                    "required": ["type", "id"],
                    "properties": {
                        "type": {
                            "type": "string",
                            "enum": ["dataset", "workbook_group", "query_result"],
                        },
                        "id": {"type": "string", "minLength": 1},
                    },
                },
                "query_result": {
                    "type": "object",
                    "additionalProperties": False,
                    "required": ["rows"],
                    "properties": {
                        "rows": {"type": "array", "items": {"type": "object"}},
                    },
                },
            },
        },
        "pipeline": {
            "type": "array",
            "minItems": 1,
            "items": {
                "type": "object",
                "additionalProperties": False,
                "required": ["stage", "config"],
                "properties": {
                    "stage": {
                        "type": "string",
                        "enum": [
                            "normalize",
                            "validate",
                            "clean",
                            "enrich",
                            "classify",
                            "score",
                            "aggregate",
                            "rank",
                            "metadata",
                            "narrate",
                            "postprocess",
                        ],
                    },
                    "config": {"type": "object"},
                },
            },
        },
        "rulesets": {
            "type": "object",
            "additionalProperties": False,
            "properties": {
                "cleaning_policy_id": {"type": "string"},
                "classification_policy_id": {"type": "string"},
                "scoring_policy_id": {"type": "string"},
            },
        },
        "outputs": {
            "type": "object",
            "additionalProperties": False,
            "required": ["logical_artifacts", "render"],
            "properties": {
                "logical_artifacts": {"type": "array"},
                "render": {
                    "type": "array",
                    "items": {
                        "type": "object",
                        "additionalProperties": False,
                        "required": ["format"],
                        "properties": {
                            "format": {"type": "string", "enum": ["json", "xlsx"]},
                            "template_id": {"type": "string"},
                            "options": {
                                "type": "object",
                                "properties": {
                                    "sheet_plan_id": {"type": "string"},
                                    "sheet_plan": {
                                        "type": "object",
                                        "additionalProperties": False,
                                        "properties": {
                                            "sheets": {
                                                "type": "array",
                                                "items": {
                                                    "type": "object",
                                                    "additionalProperties": False,
                                                    "required": ["name", "source_artifact"],
                                                    "properties": {
                                                        "name": {"type": "string", "minLength": 1},
                                                        "source_artifact": {
                                                            "type": "string",
                                                            "minLength": 1,
                                                        },
                                                        "columns": {
                                                            "type": "array",
                                                            "items": {"type": "string"},
                                                        },
                                                        "sort": {
                                                            "type": "object",
                                                            "additionalProperties": False,
                                                            "properties": {
                                                                "by": {"type": "string"},
                                                                "order": {
                                                                    "type": "string",
                                                                    "enum": ["asc", "desc"],
                                                                },
                                                            },
                                                        },
                                                        "limit": {
                                                            "type": "integer",
                                                            "minimum": 1,
                                                        },
                                                        "freeze_header": {"type": "boolean"},
                                                        "title": {"type": "string"},
                                                    },
                                                },
                                            },
                                        },
                                    },
                                },
                                "additionalProperties": True,
                            },
                        },
                    },
                },
            },
        },
        "constraints": {
            "type": "object",
            "additionalProperties": False,
            "required": ["no_fabrication", "evidence_required"],
            "properties": {
                "no_fabrication": {"type": "boolean", "const": True},
                "evidence_required": {"type": "boolean", "const": True},
                "max_output_rows": {"type": "integer", "minimum": 1},
            },
        },
    },
}


class JobSpecValidationError(ValueError):
    """Raised when an incoming job spec fails schema/semantic validation."""


@dataclass(frozen=True)
class RulesetDefinition:
    policy_id: str
    version: str
    rules: dict[str, Any]


@dataclass
class RuleTrace:
    rule_id: str
    matched_rows: int
    rejected_rows: int
    evidence_samples: list[dict[str, Any]] = field(default_factory=list)

    def to_payload(self) -> dict[str, Any]:
        return {
            "rule_id": self.rule_id,
            "matched_rows": self.matched_rows,
            "rejected_rows": self.rejected_rows,
            "evidence_samples": self.evidence_samples,
        }


@dataclass
class TransformationState:
    scope: dict[str, Any]
    records: list[dict[str, Any]] = field(default_factory=list)
    artifacts: list[dict[str, Any]] = field(default_factory=list)
    narrative: list[dict[str, Any]] = field(default_factory=list)
    validation: dict[str, Any] = field(default_factory=dict)
    metadata: dict[str, Any] = field(default_factory=dict)
    rulesets: dict[str, RulesetDefinition] = field(default_factory=dict)
    constraints: dict[str, Any] = field(default_factory=dict)


@dataclass(frozen=True)
class WorkbookSheetPlan:
    name: str
    rows: list[dict[str, Any]]
    columns: list[str]
    freeze_header: bool
    title: str | None = None


class SpecValidatorProtocol(Protocol):
    def validate(self, job_spec: dict[str, Any]) -> None:
        """Validate the incoming job spec."""


class RulesetRegistryProtocol(Protocol):
    def resolve(self, *, rulesets: dict[str, Any]) -> dict[str, RulesetDefinition]:
        """Resolve policy identifiers into concrete rulesets."""


class StageProcessorProtocol(Protocol):
    stage_name: str

    def process(
        self,
        *,
        state: TransformationState,
        config: dict[str, Any],
    ) -> tuple[TransformationState, list[RuleTrace]]:
        """Apply one stage and return updated state plus optional rule trace."""


class NarrativeComposerProtocol(Protocol):
    def compose(self, *, state: TransformationState) -> list[dict[str, Any]]:
        """Build deterministic narrative summaries."""


class RendererProtocol(Protocol):
    artifact_type: str

    def render(
        self,
        *,
        logical_result: dict[str, Any],
        render_options: dict[str, Any] | None = None,
    ) -> dict[str, Any]:
        """Render logical artifacts into physical output payload."""


class RuleEvaluator:
    """Evaluate declarative predicates using a pluggable operator map."""

    def __init__(self) -> None:
        self._operators = {
            "gt": self._op_gt,
            "lt": self._op_lt,
            "between": self._op_between,
            "contains": self._op_contains,
            "and": self._op_and,
        }

    def evaluate(self, row: dict[str, Any], predicate: dict[str, Any]) -> bool:
        operator = str(predicate.get("operator") or "").strip().lower()
        fn = self._operators.get(operator)
        if fn is None:
            return False
        return bool(fn(row, predicate))

    def _coerce_number(self, value: Any) -> float | None:
        if isinstance(value, int | float):
            return float(value)
        try:
            return float(str(value).strip())
        except Exception:
            return None

    def _op_gt(self, row: dict[str, Any], predicate: dict[str, Any]) -> bool:
        field = str(predicate.get("field") or "")
        left = self._coerce_number(row.get(field))
        right = self._coerce_number(predicate.get("value"))
        return left is not None and right is not None and left > right

    def _op_lt(self, row: dict[str, Any], predicate: dict[str, Any]) -> bool:
        field = str(predicate.get("field") or "")
        left = self._coerce_number(row.get(field))
        right = self._coerce_number(predicate.get("value"))
        return left is not None and right is not None and left < right

    def _op_between(self, row: dict[str, Any], predicate: dict[str, Any]) -> bool:
        field = str(predicate.get("field") or "")
        value = self._coerce_number(row.get(field))
        min_value = self._coerce_number(predicate.get("min"))
        max_value = self._coerce_number(predicate.get("max"))
        if value is None or min_value is None or max_value is None:
            return False
        return min_value <= value <= max_value

    def _op_contains(self, row: dict[str, Any], predicate: dict[str, Any]) -> bool:
        fields = predicate.get("fields")
        if isinstance(fields, list) and fields:
            haystack = " ".join(str(row.get(str(field), "")) for field in fields).lower()
        else:
            field = str(predicate.get("field") or "")
            haystack = str(row.get(field, "")).lower()
        needle = str(predicate.get("value") or "").strip().lower()
        return bool(needle) and needle in haystack

    def _op_and(self, row: dict[str, Any], predicate: dict[str, Any]) -> bool:
        conditions = predicate.get("conditions")
        if not isinstance(conditions, list) or not conditions:
            return False
        return all(
            isinstance(condition, dict) and self.evaluate(row, condition)
            for condition in conditions
        )


class JobSpecValidator:
    """Validate Job Spec v1.0/v1.1 schemas plus semantic constraints."""

    def __init__(self) -> None:
        self._validator = Draft7Validator(_JOB_SPEC_SCHEMA_V1)

    def validate(self, job_spec: dict[str, Any]) -> None:
        errors = sorted(self._validator.iter_errors(job_spec), key=lambda error: list(error.path))
        if errors:
            path = ".".join(str(item) for item in errors[0].path) or "root"
            raise JobSpecValidationError(f"Invalid job_spec at '{path}': {errors[0].message}")

        version = str(job_spec.get("version") or "")
        pipeline = job_spec.get("pipeline", [])
        if not isinstance(pipeline, list):
            raise JobSpecValidationError("Invalid job_spec at 'pipeline': must be a list.")

        if version == "1.0":
            unsupported = {
                str(item.get("stage") or "").strip().lower()
                for item in pipeline
                if isinstance(item, dict)
            }
            if "score" in unsupported:
                raise JobSpecValidationError(
                    "Invalid job_spec: stage 'score' requires version 1.1."
                )
            if "metadata" in unsupported:
                raise JobSpecValidationError(
                    "Invalid job_spec: stage 'metadata' requires version 1.1."
                )
        self._validate_v11_semantics(job_spec)

    def _validate_v11_semantics(self, job_spec: dict[str, Any]) -> None:  # noqa: PLR0912
        version = str(job_spec.get("version") or "")
        if version != "1.1":
            return
        pipeline = job_spec.get("pipeline", [])
        for index, stage_spec in enumerate(pipeline):
            if not isinstance(stage_spec, dict):
                continue
            stage = str(stage_spec.get("stage") or "").strip().lower()
            config = stage_spec.get("config") if isinstance(stage_spec.get("config"), dict) else {}
            if stage == "classify":
                passes = config.get("passes")
                if passes is not None and not isinstance(passes, list):
                    raise JobSpecValidationError(
                        f"Invalid job_spec at 'pipeline.{index}.config.passes': must be a list."
                    )
            if stage == "score":
                rules = config.get("rules")
                if not isinstance(rules, list) or not rules:
                    message = "must be a non-empty list."
                    raise JobSpecValidationError(
                        f"Invalid job_spec at 'pipeline.{index}.config.rules': {message}"
                    )
            if stage == "rank":
                segments = config.get("segments")
                if segments is not None and not isinstance(segments, list):
                    raise JobSpecValidationError(
                        f"Invalid job_spec at 'pipeline.{index}.config.segments': must be a list."
                    )
            if stage == "metadata":
                extractors = config.get("extractors")
                if not isinstance(extractors, list) or not extractors:
                    raise JobSpecValidationError(
                        f"Invalid job_spec at 'pipeline.{index}.config.extractors': "
                        "must be a non-empty list."
                    )
                for extractor_index, extractor in enumerate(extractors):
                    if not isinstance(extractor, dict):
                        raise JobSpecValidationError(
                            f"Invalid job_spec at 'pipeline.{index}.config.extractors."
                            f"{extractor_index}': must be an object."
                        )
                    if not str(extractor.get("extractor_id") or "").strip():
                        raise JobSpecValidationError(
                            f"Invalid job_spec at 'pipeline.{index}.config.extractors."
                            f"{extractor_index}.extractor_id': required."
                        )
                    if str(extractor.get("type") or "") not in {
                        "constant",
                        "column_value",
                        "regex",
                        "aggregate",
                    }:
                        raise JobSpecValidationError(
                            f"Invalid job_spec at 'pipeline.{index}.config.extractors."
                            f"{extractor_index}.type': unsupported extractor type."
                        )
                    if not str(extractor.get("target_key") or "").strip():
                        raise JobSpecValidationError(
                            f"Invalid job_spec at 'pipeline.{index}.config.extractors."
                            f"{extractor_index}.target_key': required."
                        )

        self._validate_sheet_plan_semantics(job_spec)

    def _validate_sheet_plan_semantics(self, job_spec: dict[str, Any]) -> None:
        outputs = job_spec.get("outputs", {}) if isinstance(job_spec.get("outputs"), dict) else {}
        render_specs = outputs.get("render", []) if isinstance(outputs.get("render"), list) else []
        declared_artifacts = {
            str(item.get("name") or "").strip()
            for item in outputs.get("logical_artifacts", [])
            if isinstance(item, dict) and str(item.get("name") or "").strip()
        }
        declared_artifacts.add("records")

        for render_index, item in enumerate(render_specs):
            if not isinstance(item, dict):
                continue
            if str(item.get("format") or "").strip().lower() != "xlsx":
                continue
            options = item.get("options") if isinstance(item.get("options"), dict) else {}
            sheet_plan = (
                options.get("sheet_plan") if isinstance(options.get("sheet_plan"), dict) else {}
            )
            sheets = sheet_plan.get("sheets") if isinstance(sheet_plan.get("sheets"), list) else []
            if not sheets:
                continue

            seen_names: set[str] = set()
            for sheet_index, sheet in enumerate(sheets):
                if not isinstance(sheet, dict):
                    raise JobSpecValidationError(
                        f"Invalid job_spec at 'outputs.render.{render_index}.options.sheet_plan."
                        f"sheets.{sheet_index}': must be an object."
                    )
                name = str(sheet.get("name") or "").strip()
                if name in seen_names:
                    raise JobSpecValidationError(
                        f"Invalid job_spec: duplicated sheet name '{name}' in sheet_plan."
                    )
                seen_names.add(name)

                source_artifact = str(sheet.get("source_artifact") or "").strip()
                if source_artifact and source_artifact not in declared_artifacts:
                    raise JobSpecValidationError(
                        "Invalid job_spec: sheet_plan source_artifact '"
                        f"{source_artifact}' is not declared in outputs.logical_artifacts."
                    )


class RulesetRegistry:
    """Code-based registry for transformation rulesets resolved by id."""

    def __init__(self) -> None:
        self._cleaning: dict[str, RulesetDefinition] = {
            "default-cleaning-v1": RulesetDefinition(
                policy_id="default-cleaning-v1",
                version="1.0",
                rules={"trim_strings": True},
            ),
        }
        self._classification: dict[str, RulesetDefinition] = {
            "default-classification-v1": RulesetDefinition(
                policy_id="default-classification-v1",
                version="1.1",
                rules={"fallback_bucket": "unclassified", "single_label": True},
            ),
        }
        self._scoring: dict[str, RulesetDefinition] = {
            "default-scoring-v1": RulesetDefinition(
                policy_id="default-scoring-v1",
                version="1.1",
                rules={"default_rank_limit": 20},
            ),
        }

    def resolve(self, *, rulesets: dict[str, Any]) -> dict[str, RulesetDefinition]:
        cleaning_id = str(rulesets.get("cleaning_policy_id") or "default-cleaning-v1")
        classification_id = str(
            rulesets.get("classification_policy_id") or "default-classification-v1"
        )
        scoring_id = str(rulesets.get("scoring_policy_id") or "default-scoring-v1")
        if cleaning_id not in self._cleaning:
            raise JobSpecValidationError(f"Unknown cleaning policy '{cleaning_id}'.")
        if classification_id not in self._classification:
            raise JobSpecValidationError(f"Unknown classification policy '{classification_id}'.")
        if scoring_id not in self._scoring:
            raise JobSpecValidationError(f"Unknown scoring policy '{scoring_id}'.")
        return {
            "cleaning": self._cleaning[cleaning_id],
            "classification": self._classification[classification_id],
            "scoring": self._scoring[scoring_id],
        }


class RulesetResolver:
    """Resolve and normalize ruleset ids from job spec into runtime rulesets."""

    def __init__(self, registry: RulesetRegistryProtocol) -> None:
        self._registry = registry

    def resolve(self, job_spec: dict[str, Any]) -> dict[str, RulesetDefinition]:
        raw = job_spec.get("rulesets", {}) if isinstance(job_spec.get("rulesets"), dict) else {}
        return self._registry.resolve(rulesets=raw)


class DecisionTraceRecorder:
    """Collect structured stage/rule telemetry for observability and auditing."""

    def __init__(self) -> None:
        self._items: list[dict[str, Any]] = []

    def record(
        self,
        *,
        stage_name: str,
        state: TransformationState,
        rule_traces: list[RuleTrace] | None = None,
    ) -> None:
        entry = {
            "stage": stage_name,
            "record_count": len(state.records),
            "artifact_count": len(state.artifacts),
            "rules": [trace.to_payload() for trace in (rule_traces or [])],
        }
        self._items.append(entry)

    def export(self) -> list[dict[str, Any]]:
        return list(self._items)


class NarrativeComposer:
    """Deterministic narrative summarizer based on computed artifacts."""

    def compose(self, *, state: TransformationState) -> list[dict[str, Any]]:
        summaries = [
            {
                "headline": "Transformation summary",
                "detail": (
                    f"Processed {len(state.records)} records and produced "
                    f"{len(state.artifacts)} artifacts."
                ),
                "evidence": {
                    "record_count": len(state.records),
                    "artifact_count": len(state.artifacts),
                },
            }
        ]
        if state.artifacts:
            summaries.append(
                {
                    "headline": "Top artifact availability",
                    "detail": (
                        "Primary artifact "
                        f"'{state.artifacts[0].get('name', 'artifact')}' is available."
                    ),
                    "evidence": {"artifact_name": state.artifacts[0].get("name")},
                }
            )
        return summaries


class LLMEnhancer:
    """Optional enhancer that never blocks deterministic narrative output."""

    def enhance(
        self,
        *,
        summaries: list[dict[str, Any]],
        context: dict[str, Any],
    ) -> list[dict[str, Any]]:
        _ = context
        return summaries


class WorkbookComposer:
    """Build workbook sheet plans from logical artifacts and declarative sheet plans."""

    def compose(
        self,
        *,
        logical_result: dict[str, Any],
        sheet_plan: dict[str, Any] | None,
    ) -> list[WorkbookSheetPlan]:
        if not isinstance(sheet_plan, dict):
            return self._fallback_summary_plan(logical_result=logical_result)

        sheets = sheet_plan.get("sheets") if isinstance(sheet_plan.get("sheets"), list) else []
        if not sheets:
            return self._fallback_summary_plan(logical_result=logical_result)

        artifacts = logical_result.get("logical_artifacts", [])
        artifact_by_name: dict[str, dict[str, Any]] = {}
        if isinstance(artifacts, list):
            for artifact in artifacts:
                if not isinstance(artifact, dict):
                    continue
                name = str(artifact.get("name") or "").strip()
                if name:
                    artifact_by_name[name] = artifact

        plans: list[WorkbookSheetPlan] = []
        for sheet in sheets:
            if not isinstance(sheet, dict):
                continue
            source_artifact_name = str(sheet.get("source_artifact") or "").strip()
            artifact = artifact_by_name.get(source_artifact_name)
            if artifact is None and source_artifact_name == "records":
                artifact = {
                    "name": "records",
                    "type": "table",
                    "rows": logical_result.get("logical_artifacts", [{}])[0].get("rows", []),
                }
            if artifact is None:
                continue
            rows = artifact.get("rows", []) if isinstance(artifact.get("rows"), list) else []
            dict_rows = [row for row in rows if isinstance(row, dict)]

            ordered_rows = self._apply_sort_and_limit(dict_rows, sheet)
            columns = self._resolve_columns(ordered_rows, sheet)
            filtered_rows = [
                {column: row.get(column) for column in columns} for row in ordered_rows
            ]
            plans.append(
                WorkbookSheetPlan(
                    name=str(sheet.get("name") or "Summary"),
                    rows=filtered_rows,
                    columns=columns,
                    freeze_header=bool(sheet.get("freeze_header", True)),
                    title=(
                        str(sheet.get("title"))
                        if isinstance(sheet.get("title"), str) and str(sheet.get("title")).strip()
                        else None
                    ),
                )
            )

        return plans or self._fallback_summary_plan(logical_result=logical_result)

    def _apply_sort_and_limit(
        self,
        rows: list[dict[str, Any]],
        sheet: dict[str, Any],
    ) -> list[dict[str, Any]]:
        output = list(rows)
        sort = sheet.get("sort") if isinstance(sheet.get("sort"), dict) else {}
        sort_field = str(sort.get("by") or "").strip()
        if sort_field:
            descending = str(sort.get("order") or "desc").strip().lower() != "asc"
            output = sorted(output, key=lambda row: row.get(sort_field) or 0, reverse=descending)
        limit_raw = sheet.get("limit")
        if isinstance(limit_raw, int) and limit_raw > 0:
            output = output[:limit_raw]
        return output

    def _resolve_columns(self, rows: list[dict[str, Any]], sheet: dict[str, Any]) -> list[str]:
        configured_columns = sheet.get("columns") if isinstance(sheet.get("columns"), list) else []
        columns = [str(column) for column in configured_columns if str(column).strip()]
        if columns:
            return columns
        if not rows:
            return []
        return [str(column) for column in rows[0].keys()]

    def _fallback_summary_plan(self, *, logical_result: dict[str, Any]) -> list[WorkbookSheetPlan]:
        artifacts = logical_result.get("logical_artifacts", [])
        if not isinstance(artifacts, list) or not artifacts:
            return [
                WorkbookSheetPlan(
                    name="Summary",
                    rows=[],
                    columns=[],
                    freeze_header=True,
                )
            ]
        first = artifacts[0] if isinstance(artifacts[0], dict) else {}
        rows = first.get("rows", []) if isinstance(first.get("rows"), list) else []
        dict_rows = [row for row in rows if isinstance(row, dict)]
        columns = [str(column) for column in dict_rows[0].keys()] if dict_rows else []
        return [
            WorkbookSheetPlan(
                name="Summary",
                rows=dict_rows,
                columns=columns,
                freeze_header=True,
            )
        ]


class JsonArtifactGenerator:
    artifact_type = "json"

    def render(
        self,
        *,
        logical_result: dict[str, Any],
        render_options: dict[str, Any] | None = None,
    ) -> dict[str, Any]:
        _ = render_options
        return {
            "format": "json",
            "filename": "transformation_result.json",
            "content": logical_result,
        }


class XlsxLogicalArtifactGenerator:
    artifact_type = "xlsx"

    def __init__(self, *, workbook_composer: WorkbookComposer | None = None) -> None:
        self._workbook_composer = workbook_composer or WorkbookComposer()

    def render(
        self,
        *,
        logical_result: dict[str, Any],
        render_options: dict[str, Any] | None = None,
    ) -> dict[str, Any]:
        workbook = Workbook()
        workbook.remove(workbook.active)

        options = render_options if isinstance(render_options, dict) else {}
        sheet_plan = (
            options.get("sheet_plan") if isinstance(options.get("sheet_plan"), dict) else {}
        )
        plans = self._workbook_composer.compose(
            logical_result=logical_result,
            sheet_plan=sheet_plan,
        )
        for plan in plans:
            self._write_sheet(workbook=workbook, plan=plan)

        stream = io.BytesIO()
        workbook.save(stream)
        encoded = base64.b64encode(stream.getvalue()).decode("ascii")
        return {
            "format": "xlsx",
            "filename": "transformation_result.xlsx",
            "content_base64": encoded,
        }

    def _write_sheet(self, *, workbook: Workbook, plan: WorkbookSheetPlan) -> None:
        sheet = workbook.create_sheet(title=plan.name[:31] or "Summary")
        cursor_row = 1

        if plan.title:
            sheet.cell(row=cursor_row, column=1, value=plan.title)
            sheet.cell(row=cursor_row, column=1).font = Font(bold=True)
            cursor_row += 1

        if plan.columns:
            sheet.append(plan.columns)
            header_row = cursor_row
            cursor_row += 1
            header_font = Font(bold=True)
            for index, column in enumerate(plan.columns, start=1):
                cell = sheet.cell(row=header_row, column=index)
                cell.font = header_font
                sheet.column_dimensions[get_column_letter(index)].width = max(
                    10,
                    len(str(column)) + 4,
                )

            for row in plan.rows:
                sheet.append([row.get(column) for column in plan.columns])

            if plan.freeze_header:
                sheet.freeze_panes = f"A{header_row + 1}"


class StageRunner:
    """Execute pipeline stage specs through a processor registry."""

    def __init__(self, processors: dict[str, StageProcessorProtocol]) -> None:
        self._processors = processors

    def run(
        self,
        *,
        state: TransformationState,
        pipeline: list[dict[str, Any]],
        recorder: DecisionTraceRecorder,
    ) -> TransformationState:
        current = state
        for stage_spec in pipeline:
            stage_name = str(stage_spec.get("stage") or "").strip().lower()
            config = stage_spec.get("config") if isinstance(stage_spec.get("config"), dict) else {}
            processor = self._processors.get(stage_name)
            if processor is None:
                recorder.record(stage_name=stage_name, state=current)
                continue
            current, rule_traces = processor.process(state=current, config=config)
            recorder.record(stage_name=stage_name, state=current, rule_traces=rule_traces)
        return current


class NormalizeStageProcessor:
    stage_name = "normalize"

    def process(
        self,
        *,
        state: TransformationState,
        config: dict[str, Any],
    ) -> tuple[TransformationState, list[RuleTrace]]:
        trim_strings = bool(config.get("trim_whitespace", True))
        if not trim_strings:
            return state, []
        normalized: list[dict[str, Any]] = []
        for row in state.records:
            normalized.append(
                {
                    str(key).strip(): value.strip() if isinstance(value, str) else value
                    for key, value in row.items()
                }
            )
        state.records = normalized
        return state, []


class ValidateStageProcessor:
    stage_name = "validate"

    def process(
        self,
        *,
        state: TransformationState,
        config: dict[str, Any],
    ) -> tuple[TransformationState, list[RuleTrace]]:
        required_columns = [
            str(item) for item in config.get("required_columns", []) if str(item).strip()
        ]
        if required_columns:
            fail_on_missing = bool(config.get("fail_on_missing", False))
            existing_columns = set(state.records[0].keys()) if state.records else set()
            missing = [column for column in required_columns if column not in existing_columns]
            if missing and fail_on_missing:
                raise JobSpecValidationError(f"Missing required columns: {', '.join(missing)}")
            state.validation["missing_columns"] = missing

        strict_types = (
            config.get("strict_types") if isinstance(config.get("strict_types"), dict) else {}
        )
        if strict_types:
            violations: list[str] = []
            for column, expected in strict_types.items():
                for row in state.records:
                    value = row.get(str(column))
                    if value is None:
                        continue
                    if not self._matches_type(value=value, expected=str(expected)):
                        violations.append(f"{column}:{expected}")
                        break
            if violations:
                raise JobSpecValidationError(
                    "Type validation failed for columns: " + ", ".join(sorted(set(violations)))
                )

        null_policies = (
            config.get("null_policies") if isinstance(config.get("null_policies"), dict) else {}
        )
        disallow_null_columns = [
            str(item)
            for item in null_policies.get("disallow_null_columns", [])
            if str(item).strip()
        ]
        if disallow_null_columns:
            for row in state.records:
                for column in disallow_null_columns:
                    if row.get(column) is None:
                        raise JobSpecValidationError(
                            f"Null values are not allowed in column '{column}'."
                        )
        return state, []

    def _matches_type(self, *, value: Any, expected: str) -> bool:
        normalized = expected.strip().lower()
        if normalized == "string":
            return isinstance(value, str)
        if normalized == "number":
            return isinstance(value, int | float)
        if normalized == "integer":
            return isinstance(value, int) and not isinstance(value, bool)
        if normalized == "boolean":
            return isinstance(value, bool)
        return True


class NoOpStageProcessor:
    def __init__(self, stage_name: str) -> None:
        self.stage_name = stage_name

    def process(
        self,
        *,
        state: TransformationState,
        config: dict[str, Any],
    ) -> tuple[TransformationState, list[RuleTrace]]:
        _ = config
        return state, []


class ClassifyStageProcessor:
    stage_name = "classify"

    def __init__(self, *, evaluator: RuleEvaluator) -> None:
        self._evaluator = evaluator

    def process(
        self,
        *,
        state: TransformationState,
        config: dict[str, Any],
    ) -> tuple[TransformationState, list[RuleTrace]]:
        ruleset_fallback = state.rulesets.get("classification")
        fallback = str(
            config.get("fallback_bucket")
            or (
                ruleset_fallback.rules.get("fallback_bucket")
                if ruleset_fallback
                else "unclassified"
            )
            or "unclassified"
        )
        target_field = str(config.get("target_field") or "classification")
        confidence_field = str(config.get("confidence_field") or "classification_confidence")
        passes = config.get("passes") if isinstance(config.get("passes"), list) else []

        if not passes:
            mappings = (
                config.get("mappings", {}) if isinstance(config.get("mappings"), dict) else {}
            )
            return self._legacy_mapping_classification(
                state=state,
                target_field=target_field,
                confidence_field=confidence_field,
                fallback=fallback,
                mappings=mappings,
            )

        traces: list[RuleTrace] = []
        for row in state.records:
            row[target_field] = fallback
            row[confidence_field] = 0.0
            for pass_index, pass_spec in enumerate(passes):
                if not isinstance(pass_spec, dict):
                    continue
                pass_rules = pass_spec.get("rules", [])
                if not isinstance(pass_rules, list):
                    continue
                matched = self._apply_rule_pass(
                    row=row,
                    pass_index=pass_index,
                    rules=pass_rules,
                    target_field=target_field,
                    confidence_field=confidence_field,
                )
                if matched:
                    break

        for pass_index, pass_spec in enumerate(passes):
            if not isinstance(pass_spec, dict):
                continue
            for rule in pass_spec.get("rules", []):
                if not isinstance(rule, dict):
                    continue
                rule_id = str(
                    rule.get("rule_id")
                    or f"classify_pass_{pass_index + 1}_{rule.get('label', 'unknown')}"
                )
                matched_rows, evidence = self._collect_rule_matches(
                    state.records, target_field, rule
                )
                traces.append(
                    RuleTrace(
                        rule_id=rule_id,
                        matched_rows=matched_rows,
                        rejected_rows=max(0, len(state.records) - matched_rows),
                        evidence_samples=evidence,
                    )
                )
        return state, traces

    def _legacy_mapping_classification(
        self,
        *,
        state: TransformationState,
        target_field: str,
        confidence_field: str,
        fallback: str,
        mappings: dict[str, Any],
    ) -> tuple[TransformationState, list[RuleTrace]]:
        traces: list[RuleTrace] = []
        for row in state.records:
            haystack = " ".join(str(value).lower() for value in row.values())
            assigned = fallback
            for label, keywords in mappings.items():
                if not isinstance(keywords, list):
                    continue
                normalized_keywords = [
                    str(keyword).strip().lower() for keyword in keywords if str(keyword).strip()
                ]
                if any(keyword in haystack for keyword in normalized_keywords):
                    assigned = str(label)
                    break
            row[target_field] = assigned
            row[confidence_field] = 1.0 if assigned != fallback else 0.0

        for label in mappings:
            matched = sum(1 for row in state.records if row.get(target_field) == label)
            traces.append(
                RuleTrace(
                    rule_id=f"legacy_mapping_{label}",
                    matched_rows=matched,
                    rejected_rows=max(0, len(state.records) - matched),
                    evidence_samples=[
                        {
                            "label": str(label),
                            "target_field": target_field,
                        }
                    ],
                )
            )
        return state, traces

    def _apply_rule_pass(
        self,
        *,
        row: dict[str, Any],
        pass_index: int,
        rules: list[Any],
        target_field: str,
        confidence_field: str,
    ) -> bool:
        _ = pass_index
        for rule in rules:
            if not isinstance(rule, dict):
                continue
            label = str(rule.get("label") or "")
            if not label:
                continue
            predicate = self._predicate_from_rule(row=row, rule=rule)
            if predicate and self._evaluator.evaluate(row, predicate):
                row[target_field] = label
                row[confidence_field] = float(rule.get("confidence") or 1.0)
                return True
        return False

    def _predicate_from_rule(
        self, *, row: dict[str, Any], rule: dict[str, Any]
    ) -> dict[str, Any] | None:
        if isinstance(rule.get("predicate"), dict):
            return rule.get("predicate")
        field = str(rule.get("field") or "")
        contains_values = rule.get("contains")
        if isinstance(contains_values, list) and contains_values:
            predicates: list[dict[str, Any]] = []
            for token in contains_values:
                token_str = str(token).strip()
                if token_str:
                    predicates.append({"operator": "contains", "field": field, "value": token_str})
            if predicates:
                return {"operator": "and", "conditions": predicates}
        keyword = str(rule.get("keyword") or "").strip()
        if keyword:
            return {"operator": "contains", "field": field, "value": keyword}

        fields = rule.get("fields")
        keywords = rule.get("keywords")
        if isinstance(fields, list) and isinstance(keywords, list):
            for keyword_item in keywords:
                keyword_str = str(keyword_item).strip()
                if keyword_str:
                    return {
                        "operator": "contains",
                        "fields": [str(item) for item in fields],
                        "value": keyword_str,
                    }
        _ = row
        return None

    def _collect_rule_matches(
        self,
        records: list[dict[str, Any]],
        target_field: str,
        rule: dict[str, Any],
    ) -> tuple[int, list[dict[str, Any]]]:
        label = str(rule.get("label") or "")
        matched_rows = [row for row in records if str(row.get(target_field) or "") == label]
        evidence = matched_rows[:2]
        return len(matched_rows), evidence


class ScoreStageProcessor:
    stage_name = "score"

    def __init__(self, *, evaluator: RuleEvaluator) -> None:
        self._evaluator = evaluator

    def process(
        self,
        *,
        state: TransformationState,
        config: dict[str, Any],
    ) -> tuple[TransformationState, list[RuleTrace]]:
        output_field = str(config.get("output_label_field") or "score_label")
        default_label = str(config.get("default_label") or "unclassified")
        rules = config.get("rules") if isinstance(config.get("rules"), list) else []

        traces: list[RuleTrace] = []
        for row in state.records:
            row[output_field] = default_label
            for rule in rules:
                if not isinstance(rule, dict):
                    continue
                label = str(rule.get("label") or "")
                when = rule.get("when") if isinstance(rule.get("when"), dict) else {}
                if label and when and self._evaluator.evaluate(row, when):
                    row[output_field] = label
                    break

        for index, rule in enumerate(rules):
            if not isinstance(rule, dict):
                continue
            label = str(rule.get("label") or "")
            matched_rows = sum(1 for row in state.records if row.get(output_field) == label)
            traces.append(
                RuleTrace(
                    rule_id=str(rule.get("rule_id") or f"score_rule_{index + 1}"),
                    matched_rows=matched_rows,
                    rejected_rows=max(0, len(state.records) - matched_rows),
                    evidence_samples=[
                        row for row in state.records if row.get(output_field) == label
                    ][:2],
                )
            )
        return state, traces


class AggregateStageProcessor:
    stage_name = "aggregate"

    def process(
        self,
        *,
        state: TransformationState,
        config: dict[str, Any],
    ) -> tuple[TransformationState, list[RuleTrace]]:
        groupings = config.get("groupings", []) if isinstance(config.get("groupings"), list) else []
        for grouping in groupings:
            if not isinstance(grouping, dict):
                continue
            name = str(grouping.get("name") or "aggregation")
            dimensions = [str(item) for item in grouping.get("dimensions", []) if str(item).strip()]
            if not dimensions:
                continue
            counts: dict[tuple[Any, ...], int] = {}
            for row in state.records:
                key = tuple(row.get(dimension) for dimension in dimensions)
                counts[key] = counts.get(key, 0) + 1
            output_rows = []
            for key, count in counts.items():
                payload = {dimensions[index]: key[index] for index in range(len(dimensions))}
                payload["count"] = count
                output_rows.append(payload)
            state.artifacts.append({"name": name, "type": "table", "rows": output_rows})
        return state, []


class RankStageProcessor:
    stage_name = "rank"

    def __init__(self, *, evaluator: RuleEvaluator) -> None:
        self._evaluator = evaluator

    def process(
        self,
        *,
        state: TransformationState,
        config: dict[str, Any],
    ) -> tuple[TransformationState, list[RuleTrace]]:
        segments = config.get("segments") if isinstance(config.get("segments"), list) else []
        if segments:
            return self._process_segmented_ranking(state=state, config=config)
        return self._process_legacy_ranking(state=state, config=config)

    def _process_legacy_ranking(
        self,
        *,
        state: TransformationState,
        config: dict[str, Any],
    ) -> tuple[TransformationState, list[RuleTrace]]:
        strategies = (
            config.get("strategies", []) if isinstance(config.get("strategies"), list) else []
        )
        if not state.artifacts or not isinstance(state.artifacts[0], dict):
            return state, []
        rows = (
            state.artifacts[0].get("rows", [])
            if isinstance(state.artifacts[0].get("rows"), list)
            else []
        )
        if not rows:
            return state, []
        for strategy in strategies:
            if not isinstance(strategy, dict):
                continue
            field = str(strategy.get("by") or "count")
            descending = str(strategy.get("order") or "desc").lower() != "asc"
            limit = int(strategy.get("limit") or 20)
            rows = sorted(rows, key=lambda row: row.get(field) or 0, reverse=descending)[:limit]
        state.artifacts[0]["rows"] = rows
        return state, []

    def _process_segmented_ranking(  # noqa: PLR0912
        self,
        *,
        state: TransformationState,
        config: dict[str, Any],
    ) -> tuple[TransformationState, list[RuleTrace]]:
        source_rows = self._resolve_source_rows(state=state, config=config)
        segments = config.get("segments") if isinstance(config.get("segments"), list) else []
        artifact_name = str(config.get("artifact_name") or "ranked")

        segment_results: dict[str, list[dict[str, Any]]] = {}
        traces: list[RuleTrace] = []
        for index, segment in enumerate(segments):
            if not isinstance(segment, dict):
                continue
            segment_name = str(segment.get("name") or f"segment_{index + 1}")
            predicate = (
                segment.get("predicate") if isinstance(segment.get("predicate"), dict) else {}
            )
            sort_by = str(segment.get("sort_by") or "count")
            order = str(segment.get("order") or "desc").strip().lower()
            descending = order != "asc"
            limit = int(segment.get("limit") or 20)

            filtered = [
                row
                for row in source_rows
                if not predicate or self._evaluator.evaluate(row, predicate)
            ]
            ranked = sorted(filtered, key=lambda row: row.get(sort_by) or 0, reverse=descending)
            ranked = ranked[:limit]
            segment_results[segment_name] = ranked

            traces.append(
                RuleTrace(
                    rule_id=str(segment.get("rule_id") or f"rank_{segment_name}"),
                    matched_rows=len(filtered),
                    rejected_rows=max(0, len(source_rows) - len(filtered)),
                    evidence_samples=ranked[:2],
                )
            )

        for segment in segments:
            if not isinstance(segment, dict):
                continue
            segment_name = str(segment.get("name") or "")
            if not segment_name:
                continue
            limit = int(segment.get("limit") or 20)
            pad_from = str(segment.get("pad_from_segment") or "").strip()
            rows = list(segment_results.get(segment_name) or [])
            if pad_from and len(rows) < limit:
                fallback_rows = list(segment_results.get(pad_from) or [])
                existing_signatures = {self._row_signature(row) for row in rows}
                for row in fallback_rows:
                    signature = self._row_signature(row)
                    if signature in existing_signatures:
                        continue
                    rows.append(row)
                    existing_signatures.add(signature)
                    if len(rows) >= limit:
                        break
            segment_results[segment_name] = rows[:limit]

        output_rows: list[dict[str, Any]] = []
        for segment in segments:
            if not isinstance(segment, dict):
                continue
            segment_name = str(segment.get("name") or "")
            if not segment_name:
                continue
            for row in segment_results.get(segment_name, []):
                materialized = dict(row)
                materialized["_segment"] = segment_name
                output_rows.append(materialized)

        state.artifacts.append({"name": artifact_name, "type": "table", "rows": output_rows})
        return state, traces

    def _resolve_source_rows(
        self,
        *,
        state: TransformationState,
        config: dict[str, Any],
    ) -> list[dict[str, Any]]:
        source_artifact = str(config.get("source_artifact") or "").strip()
        if source_artifact:
            for artifact in state.artifacts:
                if str(artifact.get("name") or "") == source_artifact and isinstance(
                    artifact.get("rows"),
                    list,
                ):
                    return [row for row in artifact["rows"] if isinstance(row, dict)]
        if state.artifacts and isinstance(state.artifacts[0].get("rows"), list):
            return [row for row in state.artifacts[0]["rows"] if isinstance(row, dict)]
        return list(state.records)

    def _row_signature(self, row: dict[str, Any]) -> tuple[tuple[str, str], ...]:
        return tuple(sorted((str(key), str(value)) for key, value in row.items()))


class MetadataStageProcessor:
    stage_name = "metadata"

    def __init__(self, *, evaluator: RuleEvaluator) -> None:
        self._evaluator = evaluator

    def process(
        self,
        *,
        state: TransformationState,
        config: dict[str, Any],
    ) -> tuple[TransformationState, list[RuleTrace]]:
        extractors = config.get("extractors") if isinstance(config.get("extractors"), list) else []
        if not extractors:
            return state, []

        traces: list[RuleTrace] = []
        for extractor in extractors:
            if not isinstance(extractor, dict):
                continue
            extractor_id = str(extractor.get("extractor_id") or "").strip()
            extractor_type = str(extractor.get("type") or "").strip()
            target_key = str(extractor.get("target_key") or "").strip()
            options = extractor.get("options") if isinstance(extractor.get("options"), dict) else {}
            required = bool(extractor.get("required", True))

            if not extractor_id or not target_key:
                continue

            value, matched_rows, evidence = self._run_extractor(
                state=state,
                extractor_type=extractor_type,
                options=options,
            )

            if value is not None:
                state.metadata[target_key] = value

            rejected_rows = max(0, len(state.records) - matched_rows)
            traces.append(
                RuleTrace(
                    rule_id=extractor_id,
                    matched_rows=matched_rows,
                    rejected_rows=rejected_rows,
                    evidence_samples=evidence,
                )
            )

            if self._is_evidence_required(state=state, required=required, evidence=evidence):
                raise JobSpecValidationError(
                    "Evidence is required for extractor "
                    f"'{extractor_id}' but no supporting rows were found."
                )

        return state, traces

    def _is_evidence_required(
        self,
        *,
        state: TransformationState,
        required: bool,
        evidence: list[dict[str, Any]],
    ) -> bool:
        return bool(state.constraints.get("evidence_required", False) and required and not evidence)

    def _run_extractor(
        self,
        *,
        state: TransformationState,
        extractor_type: str,
        options: dict[str, Any],
    ) -> tuple[Any, int, list[dict[str, Any]]]:
        if extractor_type == "constant":
            value = options.get("value")
            return value, len(state.records), [{"value": value}]
        if extractor_type == "column_value":
            return self._extract_column_value(state=state, options=options)
        if extractor_type == "regex":
            return self._extract_regex(state=state, options=options)
        if extractor_type == "aggregate":
            return self._extract_aggregate(state=state, options=options)
        return None, 0, []

    def _extract_column_value(
        self,
        *,
        state: TransformationState,
        options: dict[str, Any],
    ) -> tuple[Any, int, list[dict[str, Any]]]:
        column = str(options.get("column") or "").strip()
        if not column:
            return None, 0, []
        mode = str(options.get("mode") or "first").strip().lower()
        values = [row.get(column) for row in state.records if row.get(column) is not None]
        if not values:
            return None, 0, []
        if mode == "unique":
            unique_values = list({str(value): value for value in values}.values())
            if len(unique_values) > 1:
                raise JobSpecValidationError(
                    f"column_value extractor expects unique value in column '{column}'."
                )
            value = unique_values[0]
        else:
            value = values[0]
        evidence = [row for row in state.records if row.get(column) == value][:2]
        return value, len(evidence), evidence

    def _extract_regex(
        self,
        *,
        state: TransformationState,
        options: dict[str, Any],
    ) -> tuple[Any, int, list[dict[str, Any]]]:
        column = str(options.get("column") or "").strip()
        pattern = str(options.get("pattern") or "").strip()
        if not column or not pattern:
            return None, 0, []
        group = options.get("group", 1)
        compiled = re.compile(pattern)

        matched_rows: list[dict[str, Any]] = []
        extracted_value: Any = None
        for row in state.records:
            raw_value = row.get(column)
            if raw_value is None:
                continue
            match = compiled.search(str(raw_value))
            if not match:
                continue
            try:
                extracted_value = match.group(group)
            except Exception:
                extracted_value = match.group(0)
            matched_rows.append(row)
            if extracted_value is not None:
                break

        evidence = matched_rows[:2]
        return extracted_value, len(matched_rows), evidence

    def _extract_aggregate(  # noqa: PLR0911
        self,
        *,
        state: TransformationState,
        options: dict[str, Any],
    ) -> tuple[Any, int, list[dict[str, Any]]]:
        operation = str(options.get("operation") or "count").strip().lower()
        column = str(options.get("column") or "").strip()
        source_rows = self._resolve_source_rows(state=state, options=options)

        if operation == "count":
            return len(source_rows), len(source_rows), source_rows[:2]
        if not column:
            return None, 0, []

        values = [row.get(column) for row in source_rows if row.get(column) is not None]
        if not values:
            return None, 0, []

        if operation == "distinct":
            return len({str(value) for value in values}), len(values), source_rows[:2]
        if operation == "min":
            return min(values), len(values), source_rows[:2]
        if operation == "max":
            return max(values), len(values), source_rows[:2]
        return None, 0, []

    def _resolve_source_rows(
        self,
        *,
        state: TransformationState,
        options: dict[str, Any],
    ) -> list[dict[str, Any]]:
        source_artifact = str(options.get("source_artifact") or "").strip()
        if not source_artifact:
            return list(state.records)
        for artifact in state.artifacts:
            if str(artifact.get("name") or "") == source_artifact and isinstance(
                artifact.get("rows"),
                list,
            ):
                return [row for row in artifact.get("rows", []) if isinstance(row, dict)]
        return []


class NarrateStageProcessor:
    stage_name = "narrate"

    def __init__(
        self,
        *,
        composer: NarrativeComposerProtocol,
        llm_enhancer: LLMEnhancer,
    ) -> None:
        self._composer = composer
        self._llm_enhancer = llm_enhancer

    def process(
        self,
        *,
        state: TransformationState,
        config: dict[str, Any],
    ) -> tuple[TransformationState, list[RuleTrace]]:
        summaries = self._composer.compose(state=state)
        mode = str(config.get("mode") or "").strip().lower()
        if mode == "llm_with_fallback":
            try:
                summaries = self._llm_enhancer.enhance(
                    summaries=summaries,
                    context={"records": state.records, "artifacts": state.artifacts},
                )
            except Exception:
                _logger.warning("LLM enhancer failed, keeping deterministic narrative.")

        evidence_required = bool(state.constraints.get("evidence_required", False))
        if evidence_required:
            for summary in summaries:
                evidence = summary.get("evidence")
                if not isinstance(evidence, dict) or not evidence:
                    raise JobSpecValidationError(
                        "Narrative evidence is required by constraints.evidence_required=true."
                    )
        state.narrative = summaries
        return state, []


class LogicalResultBuilder:
    """Build canonical logical output payload independent of renderer format."""

    def build(self, *, state: TransformationState) -> dict[str, Any]:
        artifacts = list(state.artifacts)
        if not artifacts:
            artifacts = [{"name": "records", "type": "table", "rows": state.records}]
        return {
            "scope": state.scope,
            "metadata": dict(state.metadata),
            "logical_artifacts": artifacts,
            "narrative": list(state.narrative),
            "validation": dict(state.validation),
        }


class InputStateResolver:
    """Resolve input scope into a normalized transformation state."""

    def __init__(
        self,
        *,
        read_facade: AuthorizedDataRegistryReadFacade,
        query_executor: DatasetQueryExecutor,
    ) -> None:
        self._read_facade = read_facade
        self._query_executor = query_executor

    def resolve(
        self,
        *,
        context: ReadAuthorizationContext,
        job_spec: dict[str, Any],
    ) -> TransformationState:
        inputs = job_spec.get("inputs", {})
        scope = inputs.get("scope", {}) if isinstance(inputs, dict) else {}
        scope_type = str(scope.get("type") or "").strip().lower()
        scope_id = str(scope.get("id") or "").strip()

        if scope_type == "dataset":
            return self._dataset_state(context=context, dataset_id=scope_id, job_spec=job_spec)
        if scope_type == "workbook_group":
            return self._workbook_state(context=context, workbook_group_id=scope_id)
        if scope_type == "query_result":
            query_result = inputs.get("query_result", {}) if isinstance(inputs, dict) else {}
            rows = query_result.get("rows", []) if isinstance(query_result, dict) else []
            safe_rows = [row for row in rows if isinstance(row, dict)]
            return TransformationState(
                scope={"type": scope_type, "id": scope_id}, records=safe_rows
            )

        raise JobSpecValidationError("Unsupported inputs.scope.type.")

    def _dataset_state(
        self,
        *,
        context: ReadAuthorizationContext,
        dataset_id: str,
        job_spec: dict[str, Any],
    ) -> TransformationState:
        schema = self._read_facade.get_schema(context, dataset_id)
        schema_json = schema.schema_json if hasattr(schema, "schema_json") else []
        columns = [
            str(column.get("name"))
            for column in schema_json
            if isinstance(column, dict) and column.get("name")
        ][:10]
        if not columns:
            return TransformationState(scope={"type": "dataset", "id": dataset_id}, records=[])

        constraints = (
            job_spec.get("constraints", {}) if isinstance(job_spec.get("constraints"), dict) else {}
        )
        limit = int(constraints.get("max_output_rows") or 100)
        limit = max(1, min(limit, 100))
        column_list = ", ".join(f'"{column.replace(chr(34), chr(34) * 2)}"' for column in columns)
        query_result = self._query_executor.execute_query(
            context=context,
            dataset_id=dataset_id,
            sql=f"SELECT {column_list} FROM dataset LIMIT {limit}",
        )
        if query_result.get("success") is False:
            raise JobSpecValidationError(
                str(query_result.get("error") or "Unable to resolve dataset input.")
            )
        rows = query_result.get("rows", [])
        safe_rows: list[dict[str, Any]] = []
        for row in rows:
            if not isinstance(row, list):
                continue
            safe_rows.append(
                {str(columns[index]): row[index] for index in range(min(len(columns), len(row)))}
            )
        return TransformationState(
            scope={"type": "dataset", "id": dataset_id},
            records=safe_rows,
        )

    def _workbook_state(
        self,
        *,
        context: ReadAuthorizationContext,
        workbook_group_id: str,
    ) -> TransformationState:
        workbook_context = self._read_facade.get_workbook_context(context, workbook_group_id)
        rows: list[dict[str, Any]] = []
        for sheet in workbook_context.sheets:
            rows.append(
                {
                    "dataset_id": sheet.dataset_id,
                    "slug": sheet.slug,
                    "name": sheet.name,
                    "sheet_name": sheet.sheet_name,
                    "row_count": sheet.row_count,
                    "column_count": sheet.column_count,
                }
            )
        return TransformationState(
            scope={"type": "workbook_group", "id": workbook_group_id},
            records=rows,
        )


class TabularTransformationExecutionService:
    """Execute abstract tabular transformation jobs from Job Spec v1.0/v1.1."""

    def __init__(  # noqa: PLR0913
        self,
        *,
        read_facade: AuthorizedDataRegistryReadFacade | None = None,
        query_executor: DatasetQueryExecutor | None = None,
        spec_validator: SpecValidatorProtocol | None = None,
        ruleset_registry: RulesetRegistryProtocol | None = None,
        narrative_composer: NarrativeComposerProtocol | None = None,
        llm_enhancer: LLMEnhancer | None = None,
        renderers: dict[str, RendererProtocol] | None = None,
    ) -> None:
        self._read_facade = read_facade or AuthorizedDataRegistryReadFacade()
        self._query_executor = query_executor or DatasetQueryExecutor()
        self._spec_validator = spec_validator or JobSpecValidator()
        self._ruleset_registry = ruleset_registry or RulesetRegistry()
        self._ruleset_resolver = RulesetResolver(self._ruleset_registry)
        self._rule_evaluator = RuleEvaluator()
        self._narrative_composer = narrative_composer or NarrativeComposer()
        self._llm_enhancer = llm_enhancer or LLMEnhancer()
        self._logical_result_builder = LogicalResultBuilder()
        self._state_resolver = InputStateResolver(
            read_facade=self._read_facade,
            query_executor=self._query_executor,
        )
        self._renderers = renderers or {
            "json": JsonArtifactGenerator(),
            "xlsx": XlsxLogicalArtifactGenerator(),
        }
        self._engine_mode = self._resolve_engine_mode()
        self._legacy_runner = StageRunner(self._build_legacy_processors())
        self._rules_runner = StageRunner(self._build_rules_v11_processors())

    def execute(
        self,
        *,
        context: ReadAuthorizationContext,
        job_spec: dict[str, Any],
    ) -> dict[str, Any]:
        self._spec_validator.validate(job_spec)
        version = str(job_spec.get("version") or "1.0").strip()

        mode = self._engine_mode
        if mode == "rules_v1_1" and version == "1.0":
            mode = "legacy"

        state = self._state_resolver.resolve(context=context, job_spec=job_spec)
        state.rulesets = self._ruleset_resolver.resolve(job_spec)
        state.constraints = (
            job_spec.get("constraints", {}) if isinstance(job_spec.get("constraints"), dict) else {}
        )

        decision_trace = DecisionTraceRecorder()
        pipeline = [item for item in job_spec.get("pipeline", []) if isinstance(item, dict)]
        if mode == "rules_v1_1":
            state = self._rules_runner.run(state=state, pipeline=pipeline, recorder=decision_trace)
        else:
            state = self._legacy_runner.run(state=state, pipeline=pipeline, recorder=decision_trace)

        logical_result = self._logical_result_builder.build(state=state)
        rendered_outputs = self._render_outputs(
            logical_result=logical_result,
            outputs=job_spec.get("outputs", {}),
        )
        response = {
            "success": True,
            "status": 200,
            "code": "transformation_completed",
            "logical_result": logical_result,
            "decision_trace": decision_trace.export(),
            "rendered_outputs": rendered_outputs,
            "engine_mode": mode,
        }
        return sanitize_payload(response)

    def _resolve_engine_mode(self) -> str:
        configured_mode = (
            str(
                getattr(
                    settings,
                    "RESOURCE_REGISTRY_TRANSFORMATION_ENGINE_MODE",
                    _DEFAULT_TRANSFORMATION_ENGINE_MODE,
                )
                or _DEFAULT_TRANSFORMATION_ENGINE_MODE
            )
            .strip()
            .lower()
        )
        if configured_mode not in _SUPPORTED_ENGINE_MODES:
            _logger.warning(
                "Unknown RESOURCE_REGISTRY_TRANSFORMATION_ENGINE_MODE '%s'; using %s.",
                configured_mode,
                _DEFAULT_TRANSFORMATION_ENGINE_MODE,
            )
            return _DEFAULT_TRANSFORMATION_ENGINE_MODE
        return configured_mode

    def _build_legacy_processors(self) -> dict[str, StageProcessorProtocol]:
        return {
            "normalize": NormalizeStageProcessor(),
            "validate": ValidateStageProcessor(),
            "clean": NoOpStageProcessor("clean"),
            "enrich": NoOpStageProcessor("enrich"),
            "classify": ClassifyStageProcessor(evaluator=self._rule_evaluator),
            "aggregate": AggregateStageProcessor(),
            "rank": RankStageProcessor(evaluator=self._rule_evaluator),
            "narrate": NarrateStageProcessor(
                composer=self._narrative_composer,
                llm_enhancer=self._llm_enhancer,
            ),
            "postprocess": NoOpStageProcessor("postprocess"),
        }

    def _build_rules_v11_processors(self) -> dict[str, StageProcessorProtocol]:
        processors = self._build_legacy_processors()
        processors["score"] = ScoreStageProcessor(evaluator=self._rule_evaluator)
        processors["metadata"] = MetadataStageProcessor(evaluator=self._rule_evaluator)
        return processors

    def _render_outputs(
        self,
        *,
        logical_result: dict[str, Any],
        outputs: dict[str, Any],
    ) -> list[dict[str, Any]]:
        render_specs = outputs.get("render", []) if isinstance(outputs, dict) else []
        rendered: list[dict[str, Any]] = []
        for item in render_specs:
            if not isinstance(item, dict):
                continue
            format_name = str(item.get("format") or "").strip().lower()
            renderer = self._renderers.get(format_name)
            if renderer is None:
                continue
            options = item.get("options") if isinstance(item.get("options"), dict) else {}
            rendered.append(renderer.render(logical_result=logical_result, render_options=options))
        return rendered
