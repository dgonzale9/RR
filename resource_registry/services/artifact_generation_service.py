from __future__ import annotations

import csv
import io
import json
import os
import tempfile
import uuid
from abc import ABC, abstractmethod
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from django.conf import settings
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from resource_registry.models import DatasetExportPlan, DatasetGeneratedArtifact
from resource_registry.read_auth.context import ReadAuthorizationContext
from resource_registry.read_auth.exceptions import DataRegistryAuthorizationError
from resource_registry.read_auth.permissions import ReadAuthorizationService
from resource_registry.services.dataset_query_executor import DatasetQueryExecutor
from resource_registry.services.payload_sanitizer import sanitize_payload
from resource_registry.storage.factory import DataRegistryStorageFactory

_MAX_ROWS_PER_REQUIREMENT = 100


def _serialize_export_cell(value: object) -> object:
    """Serialize structured values as JSON rather than Python representations."""
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, separators=(",", ":"))
    return value


class ArtifactTooLargeError(Exception):
    """Raised when generated artifact exceeds configured size limits."""


@dataclass(frozen=True)
class GeneratedArtifactPayload:
    filename: str
    mime_type: str
    content: bytes
    metadata: dict[str, Any]


class ArtifactGenerator(ABC):
    artifact_type: str

    @abstractmethod
    def generate(
        self, *, export_plan: dict[str, Any], data_results: dict[str, list[dict[str, Any]]]
    ) -> GeneratedArtifactPayload:
        """Generate artifact bytes from export plan + query results."""


class CSVArtifactGenerator(ArtifactGenerator):
    artifact_type = "csv"

    def generate(
        self, *, export_plan: dict[str, Any], data_results: dict[str, list[dict[str, Any]]]
    ) -> GeneratedArtifactPayload:
        requirements = export_plan.get("data_requirements", [])
        source_name = requirements[0].get("name") if requirements else None
        rows = data_results.get(source_name or "", [])

        buffer = io.StringIO()
        writer = csv.writer(buffer)
        if rows:
            headers = list(rows[0].keys())
            writer.writerow(headers)
            for row in rows:
                writer.writerow([_serialize_export_cell(row.get(header)) for header in headers])
        content = buffer.getvalue().encode("utf-8")
        return GeneratedArtifactPayload(
            filename=str(export_plan.get("filename") or "dataset_export.csv"),
            mime_type="text/csv",
            content=content,
            metadata={"rows": len(rows)},
        )


class XlsxArtifactGenerator(ArtifactGenerator):
    artifact_type = "xlsx"

    def __init__(
        self,
        *,
        max_rows_per_sheet: int | None = None,
        max_columns_per_sheet: int | None = None,
        max_workbook_bytes: int | None = None,
    ) -> None:
        self._max_rows_per_sheet = int(
            settings.RESOURCE_REGISTRY_XLSX_MAX_ROWS_PER_SHEET
            if max_rows_per_sheet is None
            else max_rows_per_sheet
        )
        self._max_columns_per_sheet = int(
            settings.RESOURCE_REGISTRY_XLSX_MAX_COLUMNS_PER_SHEET
            if max_columns_per_sheet is None
            else max_columns_per_sheet
        )
        self._max_workbook_bytes = int(
            settings.RESOURCE_REGISTRY_XLSX_MAX_WORKBOOK_BYTES
            if max_workbook_bytes is None
            else max_workbook_bytes
        )

    def generate(
        self, *, export_plan: dict[str, Any], data_results: dict[str, list[dict[str, Any]]]
    ) -> GeneratedArtifactPayload:
        sheets = export_plan.get("sheets") if isinstance(export_plan.get("sheets"), list) else []
        requirements = (
            export_plan.get("data_requirements")
            if isinstance(export_plan.get("data_requirements"), list)
            else []
        )
        requirement_columns = self._build_requirement_columns(requirements)

        workbook = Workbook()
        sheet_defs = sheets or [{"name": "Summary", "source": ""}]
        for index, sheet in enumerate(sheet_defs):
            if index == 0:
                worksheet = workbook.active
            else:
                worksheet = workbook.create_sheet()
            worksheet.title = self._safe_sheet_name(str(sheet.get("name") or f"Sheet{index + 1}"))

            source = str(sheet.get("source") or "")
            rows = data_results.get(source, [])
            headers = self._resolve_headers(
                rows=rows,
                requirement_columns=requirement_columns,
                source=source,
            )
            if not headers:
                continue
            if len(headers) > self._max_columns_per_sheet:
                raise ArtifactTooLargeError
            if len(rows) > self._max_rows_per_sheet:
                raise ArtifactTooLargeError

            worksheet.append(headers)
            for row in rows:
                worksheet.append([self._safe_cell_value(_serialize_export_cell(row.get(header))) for header in headers])
            self._apply_sheet_formatting(
                worksheet=worksheet,
                header_count=len(headers),
                row_count=len(rows),
            )

        with io.BytesIO() as stream:
            workbook.save(stream)
            content = stream.getvalue()
        if len(content) > self._max_workbook_bytes:
            raise ArtifactTooLargeError
        return GeneratedArtifactPayload(
            filename=str(export_plan.get("filename") or "dataset_export.xlsx"),
            mime_type=("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
            content=content,
            metadata={"sheet_count": len(sheets) or 1},
        )

    def _build_requirement_columns(
        self,
        requirements: list[dict[str, Any]],
    ) -> dict[str, tuple[str, ...]]:
        mapping: dict[str, tuple[str, ...]] = {}
        for requirement in requirements:
            if not isinstance(requirement, dict):
                continue
            name = str(requirement.get("name") or "").strip()
            raw_columns = requirement.get("columns")
            if not name or not isinstance(raw_columns, list):
                continue
            columns = tuple(str(column).strip() for column in raw_columns if str(column).strip())
            if columns:
                mapping[name] = columns
        return mapping

    def _resolve_headers(
        self,
        *,
        rows: list[dict[str, Any]],
        requirement_columns: dict[str, tuple[str, ...]],
        source: str,
    ) -> tuple[str, ...]:
        if rows:
            return tuple(rows[0].keys())
        return requirement_columns.get(source, tuple())

    def _safe_cell_value(self, value: Any) -> Any:
        if not isinstance(value, str):
            return value
        if value.startswith(("=", "+", "-", "@")):
            return f"'{value}"
        return value

    def _apply_sheet_formatting(
        self,
        *,
        worksheet: Any,
        header_count: int,
        row_count: int,
    ) -> None:
        if header_count <= 0:
            return
        header_font = Font(bold=True)
        for index in range(1, header_count + 1):
            cell = worksheet.cell(row=1, column=index)
            cell.font = header_font
            width = max(10, min(60, len(str(cell.value or "")) + 2))
            worksheet.column_dimensions[get_column_letter(index)].width = width
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = f"A1:{get_column_letter(header_count)}{max(row_count + 1, 1)}"

    def _safe_sheet_name(self, name: str) -> str:
        invalid = set("[]:*?/\\")
        cleaned = "".join("_" if char in invalid else char for char in name).strip()
        return (cleaned or "Sheet")[:31]


class ArtifactGenerationService:
    """Generate physical artifacts from persisted DatasetExportPlan records."""

    def __init__(
        self,
        *,
        query_executor: DatasetQueryExecutor | None = None,
        storage_factory: DataRegistryStorageFactory | None = None,
        generators: dict[str, ArtifactGenerator] | None = None,
        authorization_service: ReadAuthorizationService | None = None,
    ) -> None:
        self._query_executor = query_executor or DatasetQueryExecutor()
        self._storage_factory = storage_factory or DataRegistryStorageFactory()
        self._generators = generators or {
            CSVArtifactGenerator.artifact_type: CSVArtifactGenerator(),
            XlsxArtifactGenerator.artifact_type: XlsxArtifactGenerator(),
        }
        self._authorization_service = authorization_service or ReadAuthorizationService()

    def generate_from_export_plan(
        self,
        *,
        context: ReadAuthorizationContext,
        export_plan_id: str,
        created_by: Any = None,
    ) -> dict[str, Any]:
        export_plan_record = (
            DatasetExportPlan.objects.filter(
                tenant_id=context.tenant_id,
                id=export_plan_id,
            )
            .select_related("dataset", "version")
            .first()
        )
        if export_plan_record is None:
            raise DataRegistryAuthorizationError("Artifact not found.")
        self._authorization_service.assert_can_read_dataset(
            context,
            str(export_plan_record.dataset_id),
        )

        artifact_type = str(export_plan_record.artifact_type or "").strip().lower()
        generator = self._generators.get(artifact_type)
        if generator is None:
            return {
                "success": False,
                "status": 501,
                "code": "unsupported_artifact_type",
                "error": "The requested artifact type is not supported.",
            }

        plan_json = export_plan_record.export_plan_json or {}
        data_results = self._execute_data_requirements(
            context=context,
            dataset_id=str(export_plan_record.dataset_id),
            version_id=(
                str(export_plan_record.version_id) if export_plan_record.version_id else None
            ),
            export_plan=plan_json,
        )
        try:
            generated = generator.generate(export_plan=plan_json, data_results=data_results)
        except ArtifactTooLargeError:
            return {
                "success": False,
                "status": 413,
                "code": "artifact_too_large",
                "error": "The generated XLSX artifact exceeds the supported size limits.",
            }

        artifact_uuid = uuid.uuid4()
        destination_path = (
            f"resource_registry/generated/{context.tenant_id}/{export_plan_record.dataset_id}/"
            f"{artifact_uuid}_{Path(generated.filename).name}"
        )
        storage_provider = self._storage_factory.create_provider(None)

        with tempfile.NamedTemporaryFile(prefix="dr-artifact-", delete=False) as temp_file:
            temp_file.write(generated.content)
            temp_file.flush()
            local_path = temp_file.name

        try:
            storage_metadata = storage_provider.put(local_path, destination_path)
        finally:
            try:
                os.unlink(local_path)
            except OSError:
                pass

        artifact_record = DatasetGeneratedArtifact.objects.create(
            id=artifact_uuid,
            tenant_id=context.tenant_id,
            dataset_id=export_plan_record.dataset_id,
            version_id=export_plan_record.version_id,
            export_plan=export_plan_record,
            artifact_type=artifact_type,
            filename=Path(generated.filename).name,
            storage_uri=storage_metadata.uri,
            mime_type=generated.mime_type,
            file_size_bytes=int(storage_metadata.size_bytes or len(generated.content)),
            checksum=storage_metadata.checksum,
            created_by=self._resolve_created_by(created_by),
            metadata_json=self._sanitize_metadata(
                {
                    **generated.metadata,
                    "storage_provider": storage_metadata.provider,
                    "content_type": storage_metadata.content_type,
                },
            ),
        )

        return {
            "success": True,
            "status": 200,
            "artifact": sanitize_payload(
                {
                    "artifact_id": str(artifact_record.id),
                    "artifact_type": artifact_record.artifact_type,
                    "filename": artifact_record.filename,
                    "mime_type": artifact_record.mime_type,
                    "file_size_bytes": artifact_record.file_size_bytes,
                    "created_at": artifact_record.created_at.isoformat(),
                    "storage_uri": artifact_record.storage_uri,
                },
            ),
        }

    def _execute_data_requirements(
        self,
        *,
        context: ReadAuthorizationContext,
        dataset_id: str,
        version_id: str | None,
        export_plan: dict[str, Any],
    ) -> dict[str, list[dict[str, Any]]]:
        requirements = export_plan.get("data_requirements", [])
        if not isinstance(requirements, list):
            return {}

        output: dict[str, list[dict[str, Any]]] = {}
        for requirement in requirements:
            if not isinstance(requirement, dict):
                continue
            requirement_name = str(requirement.get("name") or "").strip()
            sql = str(requirement.get("sql") or "").strip()
            if not requirement_name or not sql:
                continue

            query_result = self._query_executor.execute_query(
                context=context,
                dataset_id=dataset_id,
                version_id=version_id,
                sql=sql,
            )
            if query_result.get("success") is False:
                raise DataRegistryAuthorizationError(
                    "Unable to generate artifact from export plan."
                )

            columns = [str(column) for column in query_result.get("columns", [])]
            rows = query_result.get("rows", [])
            row_dicts: list[dict[str, Any]] = []
            for row in rows[:_MAX_ROWS_PER_REQUIREMENT]:
                if not isinstance(row, list):
                    continue
                row_dicts.append(
                    {columns[index]: row[index] for index in range(min(len(columns), len(row)))},
                )
            output[requirement_name] = row_dicts
        return output

    def _resolve_created_by(self, user: Any):
        if user is None:
            return None
        if getattr(user, "pk", None) is None:
            return None
        return user

    def _sanitize_metadata(self, metadata: dict[str, Any]) -> dict[str, Any]:
        sanitized = sanitize_payload(metadata)
        if isinstance(sanitized, dict):
            return sanitized
        return {}
